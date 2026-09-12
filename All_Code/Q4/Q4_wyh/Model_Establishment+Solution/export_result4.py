# -*- coding: utf-8 -*-
"""export_result4.py —— 把问题 4-2 的滚动结果写入官方模板 result4-2.xlsx。

【模板口径（不可变，见 others/Q4重现Q2的口径说明.md 第 7 节）】
计划购电量工作表，令行日期为 d（报告期 2025-02-01 ~ 2025-12-31）：
    列 2..144   ← x_d[1:144]            （标签 0:10-0:20 … 23:50-0:00+1）
    列 145      ← x_{d+1}[0]            （标签 0:00-0:10+1；最后一行留空）
    列 146      ← Σ x_d[0:144]          全天购电量（自然日口径，含 x[0]）
    列 147      ← Σ_t λ_{d,t}·x_d[t]    全天购电费（λ 为当日实际波动电价）
充放电量工作表：自然日六段（0:00-4:00 … 20:00-24:00），写实际执行的充/放电量与 0:00/24:00 储电量。
紧急购电量工作表：按物理区间输出，相邻连续正值合并为一个时间段事件。

【产出】Results/Tables/result4-2.xlsx（在官方模板副本上原地填写，保留模板格式）

运行：python export_result4.py
"""
from __future__ import annotations

import pickle
import sys
import traceback
from copy import copy
from datetime import datetime, time
from pathlib import Path

import numpy as np
import openpyxl

WYH = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WYH / "Data_processing"))
sys.path.insert(0, str(WYH / "Model_Establishment+Solution"))
from q4_common import BAND_LABELS, BAND_SLICES, OUT_DATA, TABLES, T, interval_labels  # noqa: E402

TEMPLATE = WYH / "others" / "原始附件_CUMCM2026_C" / "附件" / "附件5" / "result4-2.xlsx"
OUT_XLSX = TABLES / "result4-2.xlsx"
RESULTS_PATH = OUT_DATA / "q4_2_rolling_results.pkl"   # 可由 --results 覆盖（方案C 双口径用）


def load_results() -> dict:
    p = RESULTS_PATH
    if not p.exists():
        raise FileNotFoundError(f"缺少滚动结果 {p}，请先运行 q4_2_rolling.py --mode full")
    with open(p, "rb") as fh:
        return pickle.load(fh)


def _set_style(cell, style) -> None:
    cell._style = copy(style)


def fill_plan(ws, logs: list) -> None:
    """列 2..144 ← x[1:144]；列 145 ← 次日 x[0]；列 146/147 ← 自然日量与费。"""
    for i, l in enumerate(logs):
        r = i + 2
        x = np.asarray(l["plan_x"], dtype=float)
        d = datetime_from(l["date"])
        ws.cell(r, 1).value = d
        for t in range(1, T):                      # t=1..143 → 列 2..144
            ws.cell(r, t + 1).value = float(x[t])
        nxt = logs[i + 1] if i + 1 < len(logs) else None
        if nxt is not None:
            nd = datetime_from(nxt["date"])
            if (nd - d).days != 1:
                raise ValueError(f"报告期日期不连续: {l['date']} → {nxt['date']}")
            ws.cell(r, 145).value = float(nxt["plan_x"][0])
        else:
            ws.cell(r, 145).value = None
        ws.cell(r, 146).value = float(np.sum(x))
        ws.cell(r, 147).value = float(l["cost_plan"])


def datetime_from(s: str):
    return datetime.strptime(s, "%Y-%m-%d")


def fill_storage(ws, logs: list) -> None:
    """自然日六段；0:00/24:00 储电量。扩展到全部报告日。"""
    a_style = [copy(ws.cell(2, c)._style) for c in range(1, 7)]      # 含日期/时刻的行
    b_style = [copy(ws.cell(3, c)._style) for c in range(1, 7)]      # 无日期/时刻的行
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    r = 2
    for l in logs:
        d = datetime_from(l["date"])
        # 结算模块按"原样执行日前计划"运行（c_actual ≡ plan_c，r_actual ≡ plan_r），
        # 因此储能表写实际执行轨迹；若日志中存了更细的字段则优先使用。
        c_act = np.asarray(l.get("settle_c_actual", l["plan_c"]), dtype=float)
        r_act = np.asarray(l.get("settle_r_actual", l["plan_r"]), dtype=float)
        s_act = np.asarray(l["settle_s_actual"], dtype=float)
        for k, ((a, b), label) in enumerate(zip(BAND_SLICES, BAND_LABELS)):
            style = a_style if k == 0 else b_style
            for col in range(1, 7):
                _set_style(ws.cell(r, col), style[col - 1])
            if k == 0:
                ws.cell(r, 1).value = d
                ws.cell(r, 5).value = time(0, 0)
                ws.cell(r, 6).value = float(s_act[0])
            elif k == 1:
                ws.cell(r, 5).value = "24:00"
                ws.cell(r, 6).value = float(s_act[T])
            ws.cell(r, 2).value = label
            ws.cell(r, 3).value = float(np.sum(c_act[a:b]))
            ws.cell(r, 4).value = float(np.sum(r_act[a:b]))
            r += 1


def fill_emergency(ws, logs: list) -> None:
    """连续正值区间合并为事件。"""
    labels = interval_labels()
    a_style = [copy(ws.cell(2, c)._style) for c in range(1, 4)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    r = 2
    n_events = 0
    for l in logs:
        d = datetime_from(l["date"])
        e = np.asarray(l["settle_e"], dtype=float)
        t = 0
        while t < T:
            if e[t] > 1e-9:
                start = t
                while t < T and e[t] > 1e-9:
                    t += 1
                end = t - 1
                for col in range(1, 4):
                    _set_style(ws.cell(r, col), a_style[col - 1])
                ws.cell(r, 1).value = d
                ws.cell(r, 2).value = labels[start].split("-")[0] + "-" + labels[end].split("-")[1]
                ws.cell(r, 3).value = float(np.sum(e[start:end + 1]))
                r += 1
                n_events += 1
            else:
                t += 1
    print(f"    紧急购电事件数 = {n_events}")


def export() -> Path:
    res = load_results()
    logs = res["report_logs"]
    print(f"[export] 报告期 {len(logs)} 天：{logs[0]['date']} ~ {logs[-1]['date']}")
    if len(logs) != 334:
        raise ValueError(f"报告期天数应为 334，实际 {len(logs)}")

    wb = openpyxl.load_workbook(TEMPLATE)
    for name in ("计划购电量", "充放电量", "紧急购电量"):
        if name not in wb.sheetnames:
            raise KeyError(f"模板缺少工作表 {name}")
    fill_plan(wb["计划购电量"], logs)
    fill_storage(wb["充放电量"], logs)
    fill_emergency(wb["紧急购电量"], logs)
    TABLES.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"    已写出 {OUT_XLSX}")
    return OUT_XLSX


def verify() -> None:
    """spec §7 的模板来源 / 自然日总量 / 自然日费用 / 跨行恒等式 / 聚合一致性 + 完整回读校验。"""
    res = load_results()
    logs = res["report_logs"]
    wb = openpyxl.load_workbook(OUT_XLSX)
    ws = wb["计划购电量"]

    max_src = 0.0
    max_tot = 0.0
    max_cost = 0.0
    max_ident = 0.0
    for i, l in enumerate(logs):
        r = i + 2
        x = np.asarray(l["plan_x"], dtype=float)
        cells = np.array([ws.cell(r, c).value for c in range(2, 145)], dtype=float)
        max_src = max(max_src, float(np.max(np.abs(cells - x[1:144]))))
        last = ws.cell(r, 145).value
        if i + 1 < len(logs):
            chk = float(logs[i + 1]["plan_x"][0])
            max_src = max(max_src, abs(float(last) - chk))
            exp_gap = chk - x[0]
        else:
            if last is not None:
                raise AssertionError("末行末列必须为空")
            exp_gap = -x[0]
        tot = float(ws.cell(r, 146).value)
        max_tot = max(max_tot, abs(tot - float(np.sum(x))))
        # 全天购电费 = Σ λ·x（使用该日实际波动电价）
        price = np.asarray(l["price"], dtype=float)
        cost = float(ws.cell(r, 147).value)
        max_cost = max(max_cost, abs(cost - float(np.sum(price * x))))
        vis = float(np.nansum(cells)) + (float(last) if last is not None else 0.0)
        max_ident = max(max_ident, abs((vis - tot) - exp_gap))

    ws2 = wb["充放电量"]
    chg = float(np.nansum([ws2.cell(r, 3).value or 0 for r in range(2, ws2.max_row + 1)]))
    dis = float(np.nansum([ws2.cell(r, 4).value or 0 for r in range(2, ws2.max_row + 1)]))
    tgt_c = float(sum(np.sum(l.get("settle_c_actual", l["plan_c"])) for l in logs))
    tgt_r = float(sum(np.sum(l.get("settle_r_actual", l["plan_r"])) for l in logs))
    ws3 = wb["紧急购电量"]
    emg = float(np.nansum([ws3.cell(r, 3).value or 0 for r in range(2, ws3.max_row + 1)]))
    tgt_e = float(sum(np.sum(l["settle_e"]) for l in logs))

    rows = [
        ("报告期天数=334", len(logs) == 334, f"{len(logs)}"),
        ("计划列来源（列2-144 = x[1:144]）", max_src < 1e-6, f"max|Δ|={max_src:.3e}"),
        ("全天购电量 = 自然日 Σx[0:144]", max_tot < 1e-6, f"max|Δ|={max_tot:.3e}"),
        ("全天购电费 = Σ λ·x（实际波动价）", max_cost < 1e-6, f"max|Δ|={max_cost:.3e}"),
        ("跨行恒等式（容差 1e-8）", max_ident < 1e-8, f"max|Δ|={max_ident:.3e}"),
        ("充放电量聚合一致", abs(chg - tgt_c) < 1e-6 and abs(dis - tgt_r) < 1e-6,
         f"充 {chg:,.2f}/{tgt_c:,.2f}，放 {dis:,.2f}/{tgt_r:,.2f}"),
        ("紧急购电量聚合一致", abs(emg - tgt_e) < 1e-6, f"{emg:,.2f}/{tgt_e:,.2f}"),
        ("储能表行数 = 334×6", ws2.max_row - 1 == 334 * 6, f"{ws2.max_row - 1}"),
    ]
    ok = all(r[1] for r in rows)
    lines = ["# result4-2.xlsx 导出回读校验", "",
             f"- 文件：`Results/Tables/result4-2.xlsx`", ""]
    lines += ["| 校验项 | 结果 | 实测 |", "| --- | --- | --- |"]
    for name, good, val in rows:
        lines.append(f"| {name} | {'PASS' if good else 'FAIL'} | {val} |")
    (TABLES / "q4_2_export_verify.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    for name, good, val in rows:
        print(f"    [{'PASS' if good else 'FAIL'}] {name}: {val}")
    if not ok:
        raise RuntimeError("result4-2.xlsx 回读校验未通过")


if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="导出 result4-2.xlsx（默认使用主口径结果）")
    ap.add_argument("--results", type=Path, default=OUT_DATA / "q4_2_rolling_results.pkl",
                    help="滚动结果 pkl 路径（方案C 双口径可指向 variants/<tag>_results.pkl）")
    ap.add_argument("--out", type=Path, default=TABLES / "result4-2.xlsx",
                    help="输出工作簿路径")
    ap.add_argument("--verify-only", action="store_true", help="只做回读校验，不重写工作簿")
    _a = ap.parse_args()
    RESULTS_PATH = _a.results
    OUT_XLSX = _a.out
    try:
        if not _a.verify_only:
            export()
        verify()
        print("[export] 完成")
    except Exception as exc:  # noqa: BLE001
        print("=" * 70)
        traceback.print_exc()
        print(f"[export] 失败：{exc}")
        sys.exit(1)
