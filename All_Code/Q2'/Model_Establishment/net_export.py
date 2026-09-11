# -*- coding: utf-8 -*-
"""优化版 Q2 导出 result2.xlsx 与汇总表。"""
import os, pickle
import numpy as np
import pandas as pd
import openpyxl
from datetime import datetime, time as dtime
from paths import Q2_DATA_PROCESSING, Q2_TABLES, REPO_ROOT, RESULT2_TEMPLATE

T = 144
OUT_XLSX = os.path.join(Q2_TABLES, "result2.xlsx")
DATA_TRANS = os.path.join(REPO_ROOT, "All_Code", "Data_preprocessing", "Data_transformation")


def load():
    with open(os.path.join(Q2_DATA_PROCESSING, "net_rolling_results.pkl"), "rb") as fh:
        res = pickle.load(fh)
    with open(os.path.join(Q2_DATA_PROCESSING, "q2_dataset.pkl"), "rb") as fh:
        ds = pickle.load(fh)
    tm = pd.read_csv(os.path.join(DATA_TRANS, "time_map.csv"), encoding="utf-8-sig")
    return res, ds, tm


def make_mapping(ds, tm):
    rows = []
    for t in range(1, T + 1):
        row = tm[tm["time_idx"] == t].iloc[0]
        col = 145 if t == 1 else t
        rows.append([t, row["data_label"], row["template_col_label"], col])
    df = pd.DataFrame(rows, columns=["内部时段序号", "附件时间标签(结束时刻)", "模板列标签", "模板列号"])
    df.to_excel(os.path.join(Q2_TABLES, "interval_template_mapping.xlsx"), index=False)
    return df


def build_workbook(res, ds, tm):
    logs = res["report_logs"]
    tpl = openpyxl.load_workbook(RESULT2_TEMPLATE)
    plan_tpl = tpl["计划购电量"]; chg_tpl = tpl["充放电量"]; emg_tpl = tpl["紧急购电量"]
    plan_header = [plan_tpl.cell(1, c).value for c in range(1, plan_tpl.max_column + 1)]
    chg_header = [chg_tpl.cell(1, c).value for c in range(1, chg_tpl.max_column + 1)]
    emg_header = [emg_tpl.cell(1, c).value for c in range(1, emg_tpl.max_column + 1)]
    if len(plan_header) != 147:
        raise RuntimeError(f"template cols {len(plan_header)}")
    out = openpyxl.Workbook(); out.remove(out.active)
    ws1 = out.create_sheet("计划购电量"); ws1.append(plan_header)
    period_start = {int(r.time_idx): str(r.period_start) for r in tm.itertuples()}
    period_end = {int(r.time_idx): str(r.period_end) for r in tm.itertuples()}
    dates = [datetime.strptime(l["date"], "%Y-%m-%d") for l in logs]
    for i, l in enumerate(logs):
        x = l["plan_x"]
        row = [dates[i]] + [None] * 146
        for t in range(1, T):
            row[t] = float(x[t])
        if i + 1 < len(logs):
            row[144] = float(logs[i + 1]["plan_x"][0])
        else:
            row[144] = None
        row[145] = float(np.sum(x))
        row[146] = float(l["cost_plan"])
        ws1.append(row)
    ws2 = out.create_sheet("充放电量"); ws2.append(chg_header)
    windows = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    for l in logs:
        c, r, s = l["plan_c"], l["plan_r"], l["plan_s"]
        for k, (a, b) in enumerate(windows):
            row = [None] * 6
            if k == 0:
                row[0] = datetime.strptime(l["date"], "%Y-%m-%d")
                row[4] = dtime(0, 0); row[5] = float(s[0])
            elif k == 1:
                row[4] = "24:00"; row[5] = float(s[T])
            row[1] = labels[k]; row[2] = float(np.sum(c[a:b])); row[3] = float(np.sum(r[a:b]))
            ws2.append(row)
    ws3 = out.create_sheet("紧急购电量"); ws3.append(emg_header)
    for l in logs:
        e = l["settle_e"]; d = datetime.strptime(l["date"], "%Y-%m-%d"); t = 1
        while t <= T:
            if e[t - 1] > 1e-9:
                start = t
                while t <= T and e[t - 1] > 1e-9:
                    t += 1
                end = t - 1
                ws3.append([d, f"{period_start[start]}-{period_end[end]}", float(np.sum(e[start - 1:end]))])
            else:
                t += 1
    out.save(OUT_XLSX)
    return OUT_XLSX


def make_summary_tables(res, ds):
    logs = res["report_logs"]
    total_cost = float(np.sum([l["cost_total"] for l in logs]))
    plan_cost = float(np.sum([l["cost_plan"] for l in logs]))
    emg_cost = float(np.sum([l["cost_emergency"] for l in logs]))
    rows = [["指标", "数值", "单位"],
            ["报告期天数", len(logs), "天"],
            ["总费用", total_cost, "元"],
            ["计划购电费", plan_cost, "元"],
            ["紧急购电费", emg_cost, "元"],
            ["计划购电总量", float(np.sum([np.sum(l["plan_x"]) for l in logs])), "kWh"],
            ["紧急购电量", float(np.sum([l["emergency_kwh"] for l in logs])), "kWh"],
            ["弃光量", float(np.sum([l["curtail_kwh"] for l in logs])), "kWh"],
            ["未提取计划量", float(np.sum([l["unextracted_kwh"] for l in logs])), "kWh"],
            ["紧急购电率", float(np.sum([l["emergency_kwh"] for l in logs]) / np.sum([l["load_kwh"] for l in logs])), "1"]]
    pd.DataFrame(rows).to_excel(os.path.join(Q2_TABLES, "annual_cost_summary.xlsx"), index=False, header=False)

    months = [datetime.strptime(l["date"], "%Y-%m-%d").month for l in logs]
    df = pd.DataFrame(logs)
    df["month"] = months
    g = df.groupby("month").agg(
        days=("date", "count"),
        cost_total=("cost_total", "sum"),
        plan_cost=("cost_plan", "sum"),
        emergency_cost=("cost_emergency", "sum"),
        emergency_kwh=("emergency_kwh", "sum"),
        curtail_kwh=("curtail_kwh", "sum"),
        unextracted_kwh=("unextracted_kwh", "sum"),
        load_kwh=("load_kwh", "sum"),
    ).reset_index()
    g["emergency_rate"] = g["emergency_kwh"] / g["load_kwh"]
    g.to_excel(os.path.join(Q2_TABLES, "monthly_metrics.xlsx"), index=False)
    return total_cost, plan_cost, emg_cost


def run():
    res, ds, tm = load()
    make_mapping(ds, tm)
    build_workbook(res, ds, tm)
    total, plan, emg = make_summary_tables(res, ds)
    print(f"[opt export] result2.xlsx done, total={total:.2f}, plan={plan:.2f}, emergency={emg:.2f}")

if __name__ == "__main__":
    run()
