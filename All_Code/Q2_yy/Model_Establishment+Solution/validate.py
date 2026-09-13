# -*- coding: utf-8 -*-
"""独立严格校验脚本：任何一项检查失败均返回非零退出码。"""
from pathlib import Path
from datetime import datetime
import hashlib
import pickle
import sys

import numpy as np
import pandas as pd
import openpyxl

from paths import Q2_ROOT, Q2_DATA_PROCESSING, Q2_TABLES, REPO_ROOT

T, D = 144, 365
TOL = 1e-5
ATTACH_DIR = REPO_ROOT / "Data" / "附件"


class Audit:
    def __init__(self):
        self.rows = []

    def add(self, name, ok, evidence):
        self.rows.append((name, bool(ok), str(evidence)))

    @property
    def passed(self):
        return all(x[1] for x in self.rows)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def load_pickle(path):
    try:
        with open(path, "rb") as fh:
            return pickle.load(fh)
    except TypeError as exc:
        # forecasts.pkl 由较早 pandas 生成；新版本要求把旧 pickle 中的 slice
        # 显式转换为 BlockPlacement。只修复反序列化接口，不改任何缓存数值。
        if "BlockPlacement" not in str(exc):
            raise
        import pandas.core.internals.blocks as blocks
        from pandas._libs.internals import BlockPlacement
        original_new_block = blocks.new_block

        def compatible_new_block(values, placement, *, ndim, refs=None):
            if isinstance(placement, slice):
                placement = BlockPlacement(placement)
            return original_new_block(values, placement=placement, ndim=ndim, refs=refs)

        blocks.new_block = compatible_new_block
        try:
            with open(path, "rb") as fh:
                return pickle.load(fh)
        finally:
            blocks.new_block = original_new_block


def raw_data_check(audit, ds):
    a1 = pd.read_excel(ATTACH_DIR / "附件1.xlsx", sheet_name=0)
    sheets = list(pd.read_excel(ATTACH_DIR / "附件2.xlsx", sheet_name=None).values())
    raw_price = pd.to_numeric(a1.iloc[:, 1], errors="raise").to_numpy(float)
    raw_load = sheets[0].iloc[:, 1:].apply(pd.to_numeric, errors="raise").to_numpy(float).T / 6.0
    raw_pv = sheets[1].iloc[:, 1:].apply(pd.to_numeric, errors="raise").to_numpy(float).T / 6.0
    diffs = {"price": float(np.max(np.abs(ds["price"] - raw_price[:, None]))),
             "load": float(np.max(np.abs(ds["load"] - raw_load))),
             "pv": float(np.max(np.abs(ds["pv"] - raw_pv)))}
    dates0 = pd.to_datetime(sheets[0].iloc[:, 0], errors="raise").dt.strftime("%Y-%m-%d").to_numpy()
    dates1 = pd.to_datetime(sheets[1].iloc[:, 0], errors="raise").dt.strftime("%Y-%m-%d").to_numpy()
    dates_ok = np.array_equal(dates0, ds["date_str"]) and np.array_equal(dates1, ds["date_str"])
    ok = dates_ok and max(diffs.values()) <= 1e-10
    audit.add("官方附件逐点匹配", ok,
              f"price/load/pv最大绝对差={diffs}; 日期逐日一致={dates_ok}; 仅执行kW/6=kWh")
    audit.add("官方附件文件指纹", True,
              f"附件1 sha256={sha256(ATTACH_DIR/'附件1.xlsx')}; 附件2 sha256={sha256(ATTACH_DIR/'附件2.xlsx')}")


def forecast_scenario_check(audit, ds, fc):
    valid = (fc["Lhat"].shape == (T, D) and fc["Ghat"].shape == (T, D)
             and np.all(np.isfinite(fc["Lhat"])) and np.all(np.isfinite(fc["Ghat"]))
             and np.min(fc["Lhat"]) >= 0 and np.min(fc["Ghat"]) >= 0)
    audit.add("因果预测缓存完整", valid,
              f"Lhat={fc['Lhat'].shape}, Ghat={fc['Ghat'].shape}, 声明={fc.get('causality')}")

    residual_l = ds["load"] - fc["Lhat"]
    residual_g = ds["pv"] - fc["Ghat"]
    all_ok, max_diff = True, 0.0
    caches = {}
    for m in (10, 20, 30):
        cache = load_pickle(Q2_DATA_PROCESSING / f"scenarios_M{m}.pkl")
        caches[m] = cache
        ok_meta = (cache.get("history_window_days") == 30 and
                   cache.get("sampling_rule") == "joint_daily_residual_blocks_strictly_before_decision_day")
        if cache["L_all"].shape != (D, m, T) or cache["G_all"].shape != (D, m, T) or not ok_meta:
            all_ok = False
            continue
        for i in range(D):
            days = np.arange(max(1, i - 30), i, dtype=int)
            if days.size == 0:
                exp_l = np.tile(fc["Lhat"][:, i], (m, 1))
                exp_g = np.tile(fc["Ghat"][:, i], (m, 1))
            else:
                idx = np.random.default_rng(20260101 * 1000 + i).integers(0, days.size, size=m)
                picked = days[idx]
                exp_l = np.clip(fc["Lhat"][:, i][None, :] + residual_l[:, picked].T, 0, None)
                exp_g = np.clip(fc["Ghat"][:, i][None, :] + residual_g[:, picked].T, 0, None)
            d = max(float(np.max(np.abs(exp_l - cache["L_all"][i]))),
                    float(np.max(np.abs(exp_g - cache["G_all"][i]))))
            max_diff = max(max_diff, d)
            if d > 1e-10:
                all_ok = False
                break
    nested_diff = max(
        float(np.max(np.abs(caches[10]["L_all"]-caches[20]["L_all"][:, :10]))),
        float(np.max(np.abs(caches[20]["L_all"]-caches[30]["L_all"][:, :20]))),
        float(np.max(np.abs(caches[10]["G_all"]-caches[20]["G_all"][:, :10]))),
        float(np.max(np.abs(caches[20]["G_all"]-caches[30]["G_all"][:, :20]))),
    )
    nested = nested_diff <= 1e-10
    audit.add("场景严格历史且联合抽样", all_ok and nested,
              f"逐日按前30个已结束日期重构最大差={max_diff:.3e}; "
              f"M=10/20/30嵌套最大差={nested_diff:.3e}, 通过={nested}")


def result_check(audit, ds, res):
    logs = res["report_logs"]
    expected_dates = ds["date_str"][31:].tolist()
    dates_ok = [x["date"] for x in logs] == expected_dates and len(logs) == 334
    audit.add("报告期未参与选参", dates_ok and res.get("selection_period") == "2025-01-15..2025-01-31"
              and res.get("report_period") == "2025-02-01..2025-12-31",
              f"selection={res.get('selection_period')}; report={res.get('report_period')}; n={len(logs)}")
    full_logs = res.get("full_logs", [])
    state_origin_ok = (len(full_logs) > 0 and full_logs[0]["date"] == "2025-01-01"
                       and abs(float(full_logs[0]["s0"]) - 6000.0) <= TOL)
    audit.add("题目初始SOC连续传递", state_origin_ok,
              f"first_date={full_logs[0]['date'] if full_logs else None}; "
              f"first_SOC={full_logs[0]['s0'] if full_logs else None}; origin={res.get('state_origin')}")
    tune = res["tuning"]["joint"]
    chosen = tune[tune["is_selected"] == True]
    choose_ok = (len(chosen) == 1 and int(chosen.iloc[0]["M"]) == res["m_best"]
                 and abs(float(chosen.iloc[0]["beta"]) - res["beta_best"]) <= 1e-12
                 and abs(float(chosen.iloc[0]["kappa_mult"]) * res["kappa2_base"] - res["kappa2_best"]) <= 1e-12)
    audit.add("参数选择可追溯", choose_ok,
              f"唯一选中行={len(chosen)}; M={res['m_best']}, beta={res['beta_best']}, kappa2={res['kappa2_best']}")

    maxima = {"balance": 0.0, "soc_rec": 0.0, "plan_soc_rec": 0.0,
              "data": 0.0, "cost": 0.0, "emergency5x": 0.0, "cross_day": 0.0}
    violations = {"bounds": 0, "plan_bound": 0, "plan_both": 0, "actual_over_plan": 0,
                  "negative": 0, "method": 0}
    for j, l in enumerate(logs):
        i = j + 31
        y, e, g, w = l["settle_y"], l["settle_e"], l["settle_g"], l["settle_w"]
        spill = l["settle_spill"]
        ca, ra, sa = l["settle_c_actual"], l["settle_r_actual"], l["settle_s_actual"]
        pc, pr, ps, x = l["plan_c"], l["plan_r"], l["plan_s"], l["plan_x"]
        L, G, price = l["load_actual"], l["pv_actual"], l["price"]
        maxima["data"] = max(maxima["data"], float(np.max(np.abs(L-ds["load"][:, i]))),
                             float(np.max(np.abs(G-ds["pv"][:, i]))),
                             float(np.max(np.abs(price-ds["price"][:, i]))))
        maxima["balance"] = max(maxima["balance"], float(np.max(np.abs(y+e+g+ra-L-ca-spill))))
        maxima["soc_rec"] = max(maxima["soc_rec"], float(np.max(np.abs(sa[1:]-sa[:-1]-0.9*ca+ra/0.9))))
        maxima["plan_soc_rec"] = max(maxima["plan_soc_rec"], float(np.max(np.abs(ps[1:]-ps[:-1]-0.9*pc+pr/0.9))))
        if np.min(sa) < 1200-TOL or np.max(sa) > 10800+TOL or np.max(ca) > 5000/6+TOL or np.max(ra) > 5000/6+TOL:
            violations["bounds"] += 1
        if np.min(ps) < 1200-TOL or np.max(ps) > 10800+TOL or np.max(pc) > 5000/6+TOL or np.max(pr) > 5000/6+TOL:
            violations["plan_bound"] += 1
        if np.any((pc > TOL) & (pr > TOL)):
            violations["plan_both"] += 1
        if np.max(np.abs(ca-pc)) > TOL or np.max(np.abs(ra-pr)) > TOL:
            violations["actual_over_plan"] += 1
        if min(np.min(x), np.min(y), np.min(e), np.min(g), np.min(w), np.min(spill), np.min(ca), np.min(ra)) < -TOL:
            violations["negative"] += 1
        if l.get("settlement_method") != "causal_sequential_exact_plan":
            violations["method"] += 1
        cplan = float(np.sum(price*x)); cemg = float(np.sum(5*price*e))
        maxima["cost"] = max(maxima["cost"], abs(cplan-l["cost_plan"]), abs(cplan+cemg-l["cost_total"]))
        maxima["emergency5x"] = max(maxima["emergency5x"], abs(cemg-l["cost_emergency"]))
        if j and abs(logs[j-1]["final_soc"]-l["s0"]) > maxima["cross_day"]:
            maxima["cross_day"] = abs(logs[j-1]["final_soc"]-l["s0"])
    total_spill = float(sum(np.sum(x["settle_spill"]) for x in logs))
    physical_ok = (max(maxima.values()) <= TOL and all(v == 0 for v in violations.values()))
    audit.add("逐10分钟物理与执行约束", physical_ok,
              f"最大残差={maxima}; 违规计数={violations}; 供给过剩弃电={total_spill:.6f}kWh")
    terminal_ok = abs(float(logs[-1]["final_soc"]) - 6000.0) <= TOL
    audit.add("年末SOC硬约束", terminal_ok,
              f"final_SOC={float(logs[-1]['final_soc']):.9f}; constraint={res.get('terminal_constraint')}")

    total = float(sum(x["cost_total"] for x in logs))
    plan = float(sum(x["cost_plan"] for x in logs))
    emergency = float(sum(x["cost_emergency"] for x in logs))
    ekwh = float(sum(np.sum(x["settle_e"]) for x in logs))
    s0, send = float(logs[0]["s0"]), float(logs[-1]["final_soc"])
    terminal_adj = float(res["kappa2_base"] * max(0.0, s0-send))
    baseline_path = Q2_ROOT.parent / "Q2" / "Results" / "Tables" / "daily_rolling_log.csv"
    base = pd.read_csv(baseline_path, encoding="utf-8-sig")
    base_total = float(base["cost_total"].sum()); base_emergency = float(base["cost_emergency"].sum())
    target_ok = (np.isfinite(total) and np.isfinite(emergency) and total > 0
                 and abs(total-plan-emergency) <= TOL)
    audit.add("费用恒等式与原Q2对照", target_ok,
              f"total={total:.2f}, emergency={emergency:.2f}; baseline total={base_total:.2f}, emergency={base_emergency:.2f}")
    summary = pd.DataFrame([{
        "report_days": len(logs), "plan_cost_yuan": plan, "emergency_cost_yuan": emergency,
        "total_cost_yuan": total, "emergency_kwh": ekwh,
        "report_initial_soc_kwh": s0, "report_final_soc_kwh": send,
        "terminal_value_adjustment_yuan": terminal_adj,
        "terminal_adjusted_total_cost_yuan": total+terminal_adj,
        "baseline_total_cost_yuan": base_total, "baseline_emergency_cost_yuan": base_emergency,
        "total_saving_yuan": base_total-total, "emergency_saving_yuan": base_emergency-emergency,
    }])
    return summary


def workbook_check(audit, res):
    logs = res["report_logs"]
    wb = openpyxl.load_workbook(Q2_TABLES / "result2.xlsx", data_only=True)
    ws1, ws2, ws3 = wb["计划购电量"], wb["充放电量"], wb["紧急购电量"]
    plan_book = sum((ws1.cell(r, 146).value or 0) for r in range(2, ws1.max_row+1))
    cost_book = sum((ws1.cell(r, 147).value or 0) for r in range(2, ws1.max_row+1))
    c_book = sum((ws2.cell(r, 3).value or 0) for r in range(2, ws2.max_row+1))
    r_book = sum((ws2.cell(r, 4).value or 0) for r in range(2, ws2.max_row+1))
    e_book = sum((ws3.cell(r, 3).value or 0) for r in range(2, ws3.max_row+1))
    expected = [sum(np.sum(l[k]) for l in logs) for k in
                ("plan_x", "settle_c_actual", "settle_r_actual", "settle_e")]
    diffs = [abs(plan_book-expected[0]), abs(c_book-expected[1]),
             abs(r_book-expected[2]), abs(e_book-expected[3]),
             abs(cost_book-sum(l["cost_plan"] for l in logs))]
    ok = ws1.max_row == 335 and ws2.max_row == 2005 and max(diffs) <= TOL and ws1.cell(ws1.max_row, 145).value is None
    audit.add("result2提交表独立复算", ok,
              f"rows(plan/charge)={ws1.max_row-1}/{ws2.max_row-1}; 最大聚合差={max(diffs):.3e}; 12-31末列为空={ws1.cell(ws1.max_row,145).value is None}")

    wanted = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
    chosen_logs = [next(x for x in logs if x["date"] == d) for d in wanted]
    spec = pd.read_excel(Q2_TABLES / "specified_dates_results.xlsx", sheet_name=None)
    p1, p2, p3 = spec["表1购电"], spec["表2储能"], spec["表3紧急购电"]
    p1_dates = pd.to_datetime(p1["日期"]).dt.strftime("%Y-%m-%d").tolist()
    expected_c = sum(np.sum(x["settle_c_actual"]) for x in chosen_logs)
    expected_r = sum(np.sum(x["settle_r_actual"]) for x in chosen_logs)
    expected_e = sum(np.sum(x["settle_e"]) for x in chosen_logs)
    spec_diff = max(
        abs(float(p1["全天计划购电量(kWh)"].sum()) - sum(np.sum(x["plan_x"]) for x in chosen_logs)),
        abs(float(p1["全天计划购电费(元)"].sum()) - sum(x["cost_plan"] for x in chosen_logs)),
        abs(float(p2["实际充电量(kWh)"].sum()) - expected_c),
        abs(float(p2["实际放电量(kWh)"].sum()) - expected_r),
        abs(float(p3["紧急购电量(kWh)"].sum() if len(p3) else 0.0) - expected_e),
    )
    spec_ok = p1_dates == wanted and len(p2) == 24 and spec_diff <= TOL
    audit.add("题目四个指定日期直接作答", spec_ok,
              f"dates={p1_dates}; 储能汇总行={len(p2)}; 最大复算差={spec_diff:.3e}")


def run_validation():
    audit = Audit()
    ds = load_pickle(Q2_DATA_PROCESSING / "q2_dataset.pkl")
    fc = load_pickle(Q2_DATA_PROCESSING / "forecasts.pkl")
    res = load_pickle(Q2_DATA_PROCESSING / "q2_rolling_results.pkl")
    raw_data_check(audit, ds)
    forecast_scenario_check(audit, ds, fc)
    summary = result_check(audit, ds, res)
    workbook_check(audit, res)
    summary.to_csv(Q2_TABLES / "validated_cost_summary.csv", index=False, encoding="utf-8-sig")
    lines = [f"Q2_yy严格校验报告（{datetime.now():%Y-%m-%d %H:%M:%S}）",
             f"总体结果: {'全部PASS' if audit.passed else '存在FAIL'}", "-"*80]
    lines += [f"[{'PASS' if ok else 'FAIL'}] {name}: {evidence}" for name, ok, evidence in audit.rows]
    (Q2_TABLES / "validation_report.txt").write_text("\n".join(lines)+"\n", encoding="utf-8")
    print("\n".join(lines))
    return audit.passed


if __name__ == "__main__":
    try:
        sys.exit(0 if run_validation() else 2)
    except Exception as exc:
        print(f"validation crashed: {exc}")
        sys.exit(3)
