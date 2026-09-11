# -*- coding: utf-8 -*-
"""生成问题二论文 Word 文档（含表1、表2、表3）。"""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r"E:\2.University_materials\4.University_life\6.Study_materials\others\freshman_year(secong semester)\2026_GS\All_Code\Q2\Results\Tables"
RESULT2 = os.path.join(BASE, "result2.xlsx")
ANNUAL = os.path.join(BASE, "annual_cost_summary.xlsx")
TUNING = os.path.join(BASE, "tuning_sensitivity.xlsx")
MONTHLY = os.path.join(BASE, "monthly_metrics.xlsx")
OUT = r"D:\Personal\Desktop\问题二论文_更新_时间修正版.docx"

# ---------------- 读取结果数据 ----------------
xl = pd.ExcelFile(RESULT2)
plan = xl.parse("计划购电量")
chg = xl.parse("充放电量")
urg = xl.parse("紧急购电量")

dates = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
periods = ["10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10"]
plan["_d"] = pd.to_datetime(plan.iloc[:, 0]).dt.strftime("%Y-%m-%d")
table1 = {}
for d in dates:
    row = plan[plan["_d"] == d].iloc[0]
    table1[d] = {
        "p": [float(row[c]) for c in periods],
        "total": float(row["全天购电量"]),
        "fee": float(row["全天购电费"]),
    }

chg["_d"] = pd.to_datetime(chg["日期"]).dt.strftime("%Y-%m-%d")
chg["_d"] = chg["_d"].ffill()
windows = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
table2 = {}
for d in dates:
    sub = chg[chg["_d"] == d]
    soc0 = float(sub.iloc[0]["储电量"])
    soc24 = float(sub.iloc[1]["储电量"])
    w = {}
    for _, r in sub.iterrows():
        w[str(r["时间段"])] = (float(r["充电量"]), float(r["放电量"]))
    table2[d] = {"w": w, "soc0": soc0, "soc24": soc24}

urg["_d"] = pd.to_datetime(urg["日期"]).dt.strftime("%Y-%m-%d")
table3 = {}
for d in dates:
    sub = urg[urg["_d"] == d]
    table3[d] = [(str(r["购电时间段"]), float(r["购电量"])) for _, r in sub.iterrows()]

ann = pd.read_excel(ANNUAL, sheet_name="annual")
tune = pd.read_excel(TUNING, sheet_name=None)
mon = pd.read_excel(MONTHLY, sheet_name="monthly")

# ---------------- 文档与格式 ----------------
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.4
normal.paragraph_format.space_after = Pt(2)

def set_run(r, size=12, bold=False, cn="宋体", latin="Times New Roman", sub=False, sup=False):
    r.font.name = latin
    r._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.subscript = sub
    r.font.superscript = sup
    return r

def add_rich(p, s, size=12):
    i = 0
    while i < len(s):
        if s[i] in ("_", "^") and i + 1 < len(s):
            kind = s[i]
            j = i + 1
            if s[j] == "{":
                k = s.index("}", j)
                txt = s[j + 1:k]
                i = k + 1
            else:
                txt = s[j]
                i = j + 1
            r = p.add_run(txt)
            set_run(r, size=max(8, size - 3), sub=(kind == "_"), sup=(kind == "^"))
        else:
            j = i
            while j < len(s) and s[j] not in ("_", "^"):
                j += 1
            txt = s[i:j]
            i = j
            set_run(p.add_run(txt), size=size)
    return p

def heading(text, level=1):
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = 16
    elif level == 1:
        size = 14
    elif level == 2:
        size = 12
    else:
        size = 12
    set_run(p.add_run(text), size=size, bold=True, cn="黑体")
    if level == 0:
        p.paragraph_format.space_after = Pt(10)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p

def para(s, indent=True, align=None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    add_rich(p, s)
    return p

def math(s):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_rich(p, s)
    return p

def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), size=10.5, bold=True, cn="黑体")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    return p

def cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold)
    return cell

def cell_rich(cell, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    add_rich(p, text, size=size)
    return cell

# ---------------- 标题 ----------------
heading("问题二：考虑实际负荷与光伏波动的微网购电与储能调度策略", 0)

# 一、建模思路
heading("一、问题重述与建模思路", 1)
para("问题二在问题一的基础上放宽“每天电价和小区负载相同”的假设：每天电价仍相同（采用附件 1 的固定日内电价），但小区负载与光伏发电功率逐日、逐时段变化（附件 2）。微网提供的电能不得低于小区负载；若实际供给低于负载，需向外网紧急购电，其价格为交易时刻电价的 5 倍；除紧急购电费用外，其余时段的购电费用均按计划购电量计算。要求在每个决策日 0:00 制定当日计划购电策略，使全年购电总费用最小。")
para("在 0:00 制定计划时，当日实际负荷与光伏尚未实现，问题二本质上是一个不确定条件下的两阶段随机决策问题。本文采用“因果预测—经验场景生成—两阶段随机线性规划—事后结算”的框架：先用仅使用历史信息的预测模型给出次日负荷与光伏预测，再以历史残差整日块联合自助抽样生成未来场景；第一阶段确定所有场景共享的计划购电量与储能充放电计划，第二阶段针对各场景确定实际提取、紧急购电、光伏消纳与弃光量；实际数据实现后，用结算模型按真实负荷与光伏修正当日运行并统计费用。为避免对极端场景的尾部风险估计不足，目标函数引入条件风险价值 CVaR，并辅以储能软终端惩罚维持跨日电量平衡。")

# 二、符号说明
heading("二、符号说明", 1)
symbols = [
    ("π_{t}", "第 t 时段电价（元/kWh）"),
    ("L_{i,t}, G_{i,t}", "第 i 天第 t 时段的实际负荷、光伏电量（kWh）"),
    ("x_{i,t}", "计划购电量（kWh）"),
    ("c_{i,t}, r_{i,t}", "计划充电量、计划放电量（kWh）"),
    ("s_{i,t}", "储能设备电量（kWh）"),
    ("y, e, g, w, u", "场景变量：实际提取、紧急购电、光伏消纳、弃光、未提取量"),
    ("q_{ω}, ζ", "CVaR 辅助变量与分位数变量（元）"),
    ("α, β", "CVaR 置信水平与风险权重"),
    ("η_{c}, η_{r}", "充、放电效率（均取 0.9）"),
    ("κ_{2}", "储能软终端惩罚系数（元/kWh）"),
    ("M", "经验场景数量"),
]
st = doc.add_table(rows=len(symbols), cols=2)
st.style = "Table Grid"
for i, (a, b) in enumerate(symbols):
    cell_rich(st.rows[i].cells[0], a, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(st.rows[i].cells[1], b, align=WD_ALIGN_PARAGRAPH.LEFT)
st.columns[0].width = Pt(120)
st.columns[1].width = Pt(330)

# 三、模型建立
heading("三、模型建立", 1)
heading("3.1 数据准备与软终端系数校准", 2)
para("数据直接复用第 6 章全局预处理与转换的产物：附件 1 提供固定日内电价 π_{t}（144 个时段，全年复用）；附件 2 的负荷与光伏已由功率换算为时段电量，单位为 kWh。负荷电量范围为 [332.620, 1329.814] kWh，光伏电量范围为 [0, 1702.700] kWh，电价范围为 [0.3713, 1.3952] 元/kWh。")
para("本文全篇时间标签采用真实时段口径：标签“12:00—12:10”表示真实时段 12:00—12:10（即附件中时间标签为 12:10 的行/列）。官方结果模板的列名存在循环排列，result2.xlsx 按真实时段填入：第 2～144 列依次为本日真实时段 0:10—0:20 至 23:50—24:00，末列“0:00—0:10+1”为次日 0:00—0:10；2025-02-01 的首个时段（0:00—0:10）并入当日充放电量 0:00—4:00 合计，2025-12-31 行末列留空（对应 2026-01-01，超出附件数据范围）。")
para("储能软终端系数 κ_{2} 用于惩罚当日末电量偏离期初电量，以抑制储能电量的长期漂移。取问题一导出的储能边际价值序列的中位数作为基准：κ_{2,base} = 0.531667（样本量 144），并经调参最终取 κ_{2} = 0.132917。")

heading("3.2 因果预测模型", 2)
para("对负荷与光伏分别建立预测模型，特征包括：时段序号 t、星期、年内日序、前 1 天同时刻实际值、前 7 天同时刻实际值及其 7 天滚动均值、滚动标准差。候选模型为季节朴素、Ridge 回归与梯度提升树 GBM；逐日以前一日作为验证集，选择加权绝对百分比误差 WAPE 最小者作为当日预测模型。")
math("季节朴素：  ŷ_{i,t} = median_{j<i} y_{j,t}")
math("Ridge：  min_{w} ‖y − Xw‖^{2} + λ‖w‖^{2},  λ = 1.0")
math("WAPE = Σ_{t}|ŷ_{i,t} − y_{i,t}| / Σ_{t}|y_{i,t}|")

heading("3.3 经验场景生成", 2)
para("定义负荷、光伏的预测残差 ε^{L}_{j,t} = L_{j,t} − L̂_{j,t}，ε^{G}_{j,t} = G_{j,t} − Ĝ_{j,t}。以整日为块，仅对决策日之前的历史残差进行同日成对的自助抽样，将残差块叠加到当日预测上并截断非负，生成场景：")
math("L^{ω}_{i,t} = max(0, L̂_{i,t} + ε^{L}_{b,t}),   G^{ω}_{i,t} = max(0, Ĝ_{i,t} + ε^{G}_{b,t})")
para("最终场景数取 M = 30，随机种子为 20260101；负荷与光伏残差同日成对抽取，保证二者相关性不被破坏。")

heading("3.4 两阶段随机线性规划模型", 2)
para("第一阶段变量 x, c, r, s, ξ^{+}, ξ^{-} 跨场景共享，体现 0:00 决策的非预期性；第二阶段变量 y, e, g, w, u, q 按场景独立。场景 ω 的成本为")
math("C^{ω}_{i} = Σ_{t} π_{t} x_{i,t} + 5 Σ_{t} π_{t} e^{ω}_{i,t}")
para("引入 CVaR 刻画最坏 10% 场景的平均成本（α = 0.90）：")
math("q_{ω} ≥ C^{ω}_{i} − ζ,   q_{ω} ≥ 0;   CVaR = ζ + 1/((1−α)M) Σ_{ω} q_{ω}")
para("目标函数为")
math("min J_{i} = (1−β)(1/M)Σ_{ω} C^{ω}_{i} + β CVaR + ε Σ_{t}(c_{i,t}+r_{i,t}) + κ_{2}(ξ^{+}+ξ^{-})")
para("约束条件包括电量平衡、未提取量、弃光量、储能电量递推、软终端以及各类边界：")
math("y^{ω}_{t} + e^{ω}_{t} + g^{ω}_{t} + r_{t} = L^{ω}_{t} + c_{t}")
math("x_{t} − y^{ω}_{t} − u^{ω}_{t} = 0,   g^{ω}_{t} + w^{ω}_{t} = G^{ω}_{t}")
math("s_{t} − s_{t−1} − η_{c} c_{t} + r_{t}/η_{r} = 0,   s_{T} − s_{ref} = ξ^{+} − ξ^{-},  s_{ref} = s_{0}")
math("0 ≤ y^{ω}_{t} ≤ x_{t},  0 ≤ c_{t} ≤ 833.333,  0 ≤ r_{t} ≤ 833.333,  1200 ≤ s_{t} ≤ 10800,  0 ≤ g^{ω}_{t} ≤ G^{ω}_{t}")
para("其中 ε 为求解用小量惩罚，κ_{2} 为软终端惩罚，二者只用于求解，不计入报告费用。")

heading("3.5 事后结算模型", 2)
para("当日真实负荷与光伏实现后，固定 0:00 计划 x, c, r，重新求解实际运行：实际提取 0 ≤ y ≤ x，实际充放电 0 ≤ c_{actual} ≤ c、0 ≤ r_{actual} ≤ r，电量平衡 y + e + g + r_{actual} = L + c_{actual}，储能递推使用实际充放电量。结算目标为")
math("min 5 Σ_{t} π_{t} e_{t} − M_{plan} Σ_{t}(c_{actual,t} + r_{actual,t}),  M_{plan} = 20")
para("其中 M_{plan} 仅用于在可行范围内优先执行计划，不计入报告费用。报告费用为：计划购电费 cost_{plan} = Σ_{t}π_{t}x_{t}，紧急购电费 cost_{emergency} = 5Σ_{t}π_{t}e_{t}，总费用为二者之和。")

# 四、模型求解
heading("四、模型求解", 1)
para("采用 scipy.optimize.linprog（HiGHS）以稀疏矩阵形式求解。M = 30 时单日 LP 变量规模约 22,210 个，规划阶段平衡残差最大值 1.137×10^{−13}，SOC 越界数为 0，第一阶段变量结构共享成立。滚动流程为：2025-01-01 保持 SOC = 6000 kWh 不优化；2025 年 1 月整月为预热期，只积累预测残差，不进入报告；报告期为 2025-02-01 至 2025-12-31，共 334 天；调参窗口为 2025-02-01 至 02-28，共 28 天。")
para("超参数候选为 β ∈ {0, 0.25, 0.5, 0.75, 1.0}、M ∈ {10, 20, 30}、κ_{2} 倍数 ∈ {0.25, 0.5, 1.0, 2.0}。以调参窗口真实总费用最小为准则，最终取 β = 0、M = 30、κ_{2} = 0.132917、α = 0.90。")

# 五、结果与分析
heading("五、结果与分析", 1)
heading("5.1 预测精度", 2)
para("负荷预测 MAE 为 20.902–66.913 kWh，WAPE 为 0.028793–0.092357；光伏预测 MAE 为 13.386–35.488 kWh，WAPE 为 0.050998–0.089850。逐日模型选择次数为：GBM 414 次、Ridge 168 次、季节朴素 86 次，说明非线性 GBM 在多数日期更优。")

heading("5.2 全年购电策略总览", 2)
caption("全年购电策略总览（报告期 334 天）")
annual_rows = [
    ("计划购电费（元）", "14,065,827.90"),
    ("紧急购电费（元）", "1,175,367.00"),
    ("总费用（元）", "15,241,194.90"),
    ("日均总费用（元）", "45,632.32"),
    ("计划购电总量（kWh）", "22,352,000.91"),
    ("紧急购电量（kWh）", "307,585.82"),
    ("弃光量（kWh）", "1,909,728.33"),
    ("未提取计划量（kWh）", "1,557,932.11"),
    ("紧急购电率（电量）", "0.008311"),
    ("弃光率（电量）", "0.100149"),
    ("未提取计划比例（电量）", "0.069700"),
    ("SOC 越界天数", "0"),
]
at = doc.add_table(rows=len(annual_rows), cols=2)
at.style = "Table Grid"
for i, (a, b) in enumerate(annual_rows):
    cell_text(at.rows[i].cells[0], a, align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(at.rows[i].cells[1], b, align=WD_ALIGN_PARAGRAPH.CENTER)
para("紧急购电率最高的日期为 2025-06-01，当日紧急购电率为 0.099453、紧急购电量为 12,754.77 kWh；全年 SOC 越界天数为 0，表明储能电量始终处于 [1200, 10800] kWh 的安全区间。")

heading("5.3 指定日期购电与充放电结果", 2)
para("按题目要求，给出 2025-03-20、2025-06-21、2025-09-23、2025-12-21 四个指定日期的购电量与储能充放电量，见表 1 与表 2。")

caption("表 1　微网在指定时间段的购电量及全天的购电量和购电费（单位：kWh、元）")
para("注：表 1 与表 3 中的时段均为真实时段（如 12:00—12:10 即真实 12:00—12:10）；表 2 的六段 4 小时窗口为真实时段合计。")
t2 = doc.add_table(rows=1 + len(periods) + 2, cols=5)
t2.style = "Table Grid"
cell_text(t2.rows[0].cells[0], "时间段")
for j, d in enumerate(dates):
    cell_text(t2.rows[0].cells[j + 1], d)
for i, pr in enumerate(periods):
    cell_text(t2.rows[1 + i].cells[0], pr, align=WD_ALIGN_PARAGRAPH.LEFT)
    for j, d in enumerate(dates):
        cell_text(t2.rows[1 + i].cells[j + 1], f"{table1[d]['p'][i]:.2f}")
cell_text(t2.rows[1 + len(periods)].cells[0], "全天购电量", align=WD_ALIGN_PARAGRAPH.LEFT)
for j, d in enumerate(dates):
    cell_text(t2.rows[1 + len(periods)].cells[j + 1], f"{table1[d]['total']:.2f}")
cell_text(t2.rows[2 + len(periods)].cells[0], "全天购电费", align=WD_ALIGN_PARAGRAPH.LEFT)
for j, d in enumerate(dates):
    cell_text(t2.rows[2 + len(periods)].cells[j + 1], f"{table1[d]['fee']:.2f}")

caption("表 2　储能设备在指定时间段的充放电量及 0:00 和 24:00 的储电量（单位：kWh）")
t3 = doc.add_table(rows=1 + 1 + len(windows) + 2, cols=9)
t3.style = "Table Grid"
cell_text(t3.rows[0].cells[0], "时间段")
for j, d in enumerate(dates):
    a = t3.rows[0].cells[1 + 2 * j]
    b = t3.rows[0].cells[2 + 2 * j]
    a.merge(b)
    cell_text(a, d)
cell_text(t3.rows[1].cells[0], "时间段")
for j in range(4):
    cell_text(t3.rows[1].cells[1 + 2 * j], "充电量")
    cell_text(t3.rows[1].cells[2 + 2 * j], "放电量")
for i, win in enumerate(windows):
    cell_text(t3.rows[2 + i].cells[0], win, align=WD_ALIGN_PARAGRAPH.LEFT)
    for j, d in enumerate(dates):
        c, r = table2[d]["w"][win]
        cell_text(t3.rows[2 + i].cells[1 + 2 * j], f"{c:.2f}")
        cell_text(t3.rows[2 + i].cells[2 + 2 * j], f"{r:.2f}")
cell_text(t3.rows[2 + len(windows)].cells[0], "0:00 储电量", align=WD_ALIGN_PARAGRAPH.LEFT)
for j, d in enumerate(dates):
    a = t3.rows[2 + len(windows)].cells[1 + 2 * j]
    b = t3.rows[2 + len(windows)].cells[2 + 2 * j]
    a.merge(b)
    cell_text(a, f"{table2[d]['soc0']:.2f}")
cell_text(t3.rows[3 + len(windows)].cells[0], "24:00 储电量", align=WD_ALIGN_PARAGRAPH.LEFT)
for j, d in enumerate(dates):
    a = t3.rows[3 + len(windows)].cells[1 + 2 * j]
    b = t3.rows[3 + len(windows)].cells[2 + 2 * j]
    a.merge(b)
    cell_text(a, f"{table2[d]['soc24']:.2f}")

heading("5.4 指定日期紧急购电结果", 2)
para("四个指定日期的紧急购电时段与购电量见表 3。")
caption("表 3　微网在指定日期的紧急购电量（单位：kWh）")
maxrows = max(len(table3[d]) for d in dates)
t4 = doc.add_table(rows=2 + maxrows, cols=8)
t4.style = "Table Grid"
for j, d in enumerate(dates):
    a = t4.rows[0].cells[2 * j]
    b = t4.rows[0].cells[2 * j + 1]
    a.merge(b)
    cell_text(a, d)
for j in range(4):
    cell_text(t4.rows[1].cells[2 * j], "时间段")
    cell_text(t4.rows[1].cells[2 * j + 1], "购电量")
for i in range(maxrows):
    for j, d in enumerate(dates):
        if i < len(table3[d]):
            it, am = table3[d][i]
            cell_text(t4.rows[2 + i].cells[2 * j], it, align=WD_ALIGN_PARAGRAPH.LEFT)
            cell_text(t4.rows[2 + i].cells[2 * j + 1], f"{am:.2f}")
        else:
            cell_text(t4.rows[2 + i].cells[2 * j], "")
            cell_text(t4.rows[2 + i].cells[2 * j + 1], "")

# 六、模型检验与灵敏度分析
heading("六、模型检验与灵敏度分析", 1)
para("自动校验 12 项全部通过：功率转电量乘 1/6 一致；结算平衡残差最大值 3.411×10^{−13}；SOC 始终在 [1200, 10800]；充放电不超过 833.333 kWh；跨日 SOC 传递一致；因果性成立（特征与残差库均早于决策日）；第一阶段变量跨场景共享；弃光 w = G − g 且非负；紧急购电 5 倍电价复算一致；真实时间口径模板列映射与总量一致；报告费用不含罚项；报告期样本量 334 天。")

caption("表 4　超参数调优结果（调参窗口 2025-02-01 至 02-28，共 28 天）")
beta = tune["beta"]; mdf = tune["M"]; k2 = tune["kappa2"]
t5 = doc.add_table(rows=1 + len(beta) + len(mdf) + len(k2), cols=5)
t5.style = "Table Grid"
for j, h in enumerate(["超参数", "取值", "调参窗口总费用（元）", "紧急购电率", "是否选择"]):
    cell_text(t5.rows[0].cells[j], h, bold=True)
rr = 1
for _, row in beta.iterrows():
    cell_text(t5.rows[rr].cells[0], "β", align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(t5.rows[rr].cells[1], f"{row['beta']:.2f}")
    cell_text(t5.rows[rr].cells[2], f"{row['tuning_cost_total']:,.2f}")
    cell_text(t5.rows[rr].cells[3], f"{row['tuning_emergency_rate']:.6f}")
    cell_text(t5.rows[rr].cells[4], "√" if row["is_selected"] else "")
    rr += 1
for _, row in mdf.iterrows():
    cell_text(t5.rows[rr].cells[0], "M", align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(t5.rows[rr].cells[1], f"{int(row['M'])}")
    cell_text(t5.rows[rr].cells[2], f"{row['tuning_cost_total']:,.2f}")
    cell_text(t5.rows[rr].cells[3], f"{row['tuning_emergency_rate']:.6f}")
    cell_text(t5.rows[rr].cells[4], "√" if row["is_selected"] else "")
    rr += 1
for _, row in k2.iterrows():
    cell_rich(t5.rows[rr].cells[0], "κ_{2} 倍数", align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(t5.rows[rr].cells[1], f"{row['kappa2_mult']:.2f}")
    cell_text(t5.rows[rr].cells[2], f"{row['tuning_cost_total']:,.2f}")
    cell_text(t5.rows[rr].cells[3], f"{row['tuning_emergency_rate']:.6f}")
    cell_text(t5.rows[rr].cells[4], "√" if row["is_selected"] else "")
    rr += 1

para("调参结果表明：风险权重 β 增大虽然降低了紧急购电率，但显著抬高了计划购电费，故最终取 β = 0（期望费用口径）；场景数 M 增加至 30 后总费用略有下降；κ_{2} 在 0.25–0.5 倍区间对结果影响很小，取较小值 0.132917 即可。敏感性实验“全量提取 y = x”在调参窗口不可行，说明在负荷与光伏偏离预测较大时，强制全额提取计划电量会导致储能越界或供需失衡，从而印证“允许少提计划电量”的建模必要性。")
para("综上，本文模型满足 0:00 决策的非预期性、电价 5 倍紧急购电与储能物理约束，全年总购电费用为 15,241,194.90 元，紧急购电率仅 0.8311%，弃光率 10.0149%，未提取计划比例 6.9700%，储能电量全程处于安全区间，验证了模型的有效性。")

doc.save(OUT)
print("saved:", OUT)
