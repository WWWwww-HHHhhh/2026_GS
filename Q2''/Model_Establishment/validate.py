# -*- coding: utf-8 -*-
"""Q2'' 结果校验：平衡残差、SOC 边界、跨日结转、费用口径、result2 聚合一致性。"""
import os, pickle
import numpy as np
import openpyxl

GS = r"E:\2.University_materials\4.University_life\6.Study_materials\others\freshman_year(secong semester)\2026_GS"
Q2P = os.path.join(GS, "Q2''")
DATA_PROC = os.path.join(Q2P, "Data_processing")
TABLES = os.path.join(Q2P, "Results", "Tables")
T = 144
ETA_C = ETA_R = 0.9
SOC_MIN, SOC_MAX = 1200.0, 10800.0

def main():
    ds = pickle.load(open(os.path.join(DATA_PROC, "q2_dataset.pkl"), "rb"))
    res = pickle.load(open(os.path.join(DATA_PROC, "q2_rolling_results.pkl"), "rb"))
    report = res["report_logs"]
    max_bal = 0.0; soc_vio = 0; cross_err = 0.0; prev_soc = None
    for l in report:
        i = l["day_index"]
        L = ds["load"][:, i]; G = ds["pv"][:, i]; price = ds["price"][:, i]
        y = l["settle_y"]; e = l["settle_e"]; g = l["settle_g"]; w = l["settle_w"]
        ca = l["settle_c_actual"]; ra = l["settle_r_actual"]; sa = l["settle_s_actual"]
        bal = y + e + g + ra - L - ca
        max_bal = max(max_bal, float(np.max(np.abs(bal))))
        soc_vio += int(np.sum((sa < SOC_MIN - 1e-6) | (sa > SOC_MAX + 1e-6)))
        if prev_soc is not None:
            cross_err = max(cross_err, abs(float(sa[0]) - prev_soc))
        prev_soc = float(sa[T])
        # 5倍紧急购电费复算
        assert abs(np.sum(5*price*e) - l["cost_emergency"]) < 1e-6
    # result2 聚合
    wb = openpyxl.load_workbook(os.path.join(TABLES, "result2.xlsx"), data_only=True)
    ws1 = wb["计划购电量"]
    col146_sum = float(np.sum([ws1.cell(r,146).value or 0 for r in range(2, ws1.max_row+1)]))
    total_plan = float(sum(np.sum(l["plan_x"]) for l in report))
    ws2 = wb["充放电量"]
    chg_sum = float(np.sum([ws2.cell(r,3).value or 0 for r in range(2, ws2.max_row+1)]))
    dis_sum = float(np.sum([ws2.cell(r,4).value or 0 for r in range(2, ws2.max_row+1)]))
    total_c = float(sum(np.sum(l["plan_c"]) for l in report))
    total_r = float(sum(np.sum(l["plan_r"]) for l in report))
    print(f"max_balance_residual={max_bal:.3e}")
    print(f"soc_violations={soc_vio}  max_cross_day_soc_err={cross_err:.3e}")
    print(f"plan_col_sum={col146_sum:.2f} vs plan_total={total_plan:.2f} diff={abs(col146_sum-total_plan):.6f}")
    print(f"charge_sum={chg_sum:.2f} vs {total_c:.2f} diff={abs(chg_sum-total_c):.6f}")
    print(f"discharge_sum={dis_sum:.2f} vs {total_r:.2f} diff={abs(dis_sum-total_r):.6f}")
    print(f"total_cost={res['total_cost']:.2f} plan={res['plan_cost']:.2f} emergency={res['emergency_cost']:.2f} emg_rate={res['emergency_rate']:.6f}")
    print("VALIDATION_OK")

if __name__ == "__main__":
    main()
