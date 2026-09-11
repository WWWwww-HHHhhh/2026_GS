# -*- coding: utf-8 -*-
"""Q2'' 数据准备：只读真实附件1/附件2，功率*1/6h转电量，kappa2_base 取 Q1 储能边际价值中位数。"""
import os, pickle
from datetime import datetime, timedelta
import numpy as np
import pandas as pd
import openpyxl

GS = r"E:\2.University_materials\4.University_life\6.Study_materials\others\freshman_year(secong semester)\2026_GS"
ATT1 = os.path.join(GS, "Data", "附件", "附件1.xlsx")
ATT2 = os.path.join(GS, "Data", "附件", "附件2.xlsx")
Q1_MV = os.path.join(GS, "All_Code", "Q1", "Results", "Tables", "q1_storage_marginal_value_for_q2.csv")
OUT = os.path.join(GS, "Q2''", "Data_processing")
T, D = 144, 365
DELTA_T = 1.0 / 6.0

def main():
    os.makedirs(OUT, exist_ok=True)
    wb1 = openpyxl.load_workbook(ATT1, read_only=True, data_only=True)
    ws1 = wb1["Sheet1"]
    rows1 = list(ws1.iter_rows(min_row=2, values_only=True))
    if len(rows1) != T:
        raise RuntimeError(f"附件1 时段数应为 {T}，实际 {len(rows1)}")
    price_day = np.array([float(r[1]) for r in rows1], dtype=float)
    if (price_day <= 0).any():
        raise RuntimeError("附件1 存在非正电价")
    price = np.tile(price_day[:, None], (1, D))

    wb2 = openpyxl.load_workbook(ATT2, read_only=True, data_only=True)
    def read_series(sheet_name):
        ws = wb2[sheet_name]
        data = list(ws.iter_rows(min_row=2, values_only=True))
        if len(data) != D:
            raise RuntimeError(f"{sheet_name} 行数应为 {D}，实际 {len(data)}")
        mat = np.array([[float(v) for v in row[1:145]] for row in data], dtype=float)
        return mat.T
    load = read_series("小区负载") * DELTA_T
    pv = read_series("光伏发电实际功率") * DELTA_T

    start = datetime(2025, 1, 1)
    dates = [start + timedelta(days=i) for i in range(D)]
    date_str = np.array([d.strftime("%Y-%m-%d") for d in dates], dtype=object)

    mv = pd.read_csv(Q1_MV)
    kappa2_base = float(np.median(mv["storage_marginal_value_yuan_per_kwh"].to_numpy()))

    def end_label(t):
        minutes = t * 10
        return "0:00+1" if minutes == 1440 else f"{minutes//60}:{minutes%60:02d}"
    time_mapping = {"internal_idx": np.arange(1, T + 1), "time_labels": [end_label(t) for t in range(1, T + 1)]}
    constants = {"T": T, "D": D, "DELTA_T": DELTA_T, "CHARGE_CAP_KWH": 5000.0 * DELTA_T,
                 "ETA_C": 0.9, "ETA_R": 0.9, "SOC_MIN": 1200.0, "SOC_MAX": 10800.0, "SOC0": 6000.0}
    ds = {"dates": np.array(dates, dtype=object), "date_str": date_str, "price": price,
          "load": load, "pv": pv, "time_mapping": time_mapping, "kappa2_base": kappa2_base,
          "kappa2_stats": {"median": kappa2_base}, "constants": constants}
    with open(os.path.join(OUT, "q2_dataset.pkl"), "wb") as f:
        pickle.dump(ds, f, protocol=pickle.HIGHEST_PROTOCOL)
    print(f"saved q2_dataset.pkl: load={load.shape} pv={pv.shape} price={price.shape}")
    print(f"load kWh [{load.min():.3f},{load.max():.3f}] pv kWh [{pv.min():.3f},{pv.max():.3f}] price [{price.min():.4f},{price.max():.4f}] kappa2_base={kappa2_base:.6f}")

if __name__ == "__main__":
    main()
