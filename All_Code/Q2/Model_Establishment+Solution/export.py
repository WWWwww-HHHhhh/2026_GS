# -*- coding: utf-8 -*-
"""
P4 官方模板导出 result2.xlsx（2026 数模 C 题 问题二）
====================================================
实现依据：总纲 2.1（模板表头偏移风险）、4.8（连续紧急购电时段合并）、7.2 第 12 项。
导出结构与官方模板完全一致（sheet 名、列名、列顺序、日期格式），不新增 sheet/列。
"""
import os
import pickle
import traceback
from datetime import datetime, time as dtime

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

GS = r"E:\2.University_materials\4.University_life\6.Study_materials\others\freshman_year(secong semester)\2026_GS"
Q2CODE = os.path.join(GS, r"All_Code\Q2")
DATA_PROC = os.path.join(Q2CODE, "Data_processing")
TABLES = os.path.join(Q2CODE, "Results", "Tables")
DATA_TRANS = os.path.join(GS, r"All_Code\Data_preprocessing\Data_transformation")
TEMPLATE_XLSX = r"D:\Personal\Temp\result2.xlsx"
OUT_XLSX = os.path.join(TABLES, "result2.xlsx")

T = 144


def load_inputs():
    with open(os.path.join(DATA_PROC, "q2_rolling_results.pkl"), "rb") as fh:
        res = pickle.load(fh)
    with open(os.path.join(DATA_PROC, "q2_dataset.pkl"), "rb") as fh:
        ds = pickle.load(fh)
    tm = pd.read_csv(os.path.join(DATA_TRANS, "time_map.csv"), encoding="utf-8-sig")
    return res, ds, tm


def make_interval_mapping(ds, tm):
    """内部时段序号 <-> 附件时间标签 <-> 模板列号 三列映射。"""
    rows = []
    for t in range(1, T + 1):
        row = tm[tm["time_idx"] == t].iloc[0]
        rows.append([t, row["data_label"], row["template_col_label"], t + 1])
    df = pd.DataFrame(rows, columns=["内部时段序号", "附件时间标签(结束时刻)", "模板列标签", "模板列号"])
    path = os.path.join(TABLES, "interval_template_mapping.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="mapping")
    print(f"    时段-模板映射已写入: {path}")
    return df


def parse_date(ds_str):
    return datetime.strptime(ds_str, "%Y-%m-%d")


def build_workbook(res, ds, tm):
    """按官方模板结构填充三个 sheet。"""
    report_logs = res["report_logs"]
    # 官方模板结构
    tpl = openpyxl.load_workbook(TEMPLATE_XLSX)
    plan_tpl = tpl["计划购电量"]
    chg_tpl = tpl["充放电量"]
    emg_tpl = tpl["紧急购电量"]
    plan_header = [plan_tpl.cell(1, c).value for c in range(1, plan_tpl.max_column + 1)]
    chg_header = [chg_tpl.cell(1, c).value for c in range(1, chg_tpl.max_column + 1)]
    emg_header = [emg_tpl.cell(1, c).value for c in range(1, emg_tpl.max_column + 1)]
    if len(plan_header) != 147:
        raise ValueError(f"计划购电量模板列数应为 147，实际 {len(plan_header)}")

    out = openpyxl.Workbook()
    out.remove(out.active)

    # ---- Sheet 1 计划购电量 ----
    ws1 = out.create_sheet("计划购电量")
    ws1.append(plan_header)
    period_start = {int(r.time_idx): str(r.period_start) for r in tm.itertuples()}
    period_end = {int(r.time_idx): str(r.period_end) for r in tm.itertuples()}

    for l in report_logs:
        d = parse_date(l["date"])
        x = l["plan_x"]                       # (144,)
        row = [d] + [None] * 146
        for t in range(T):
            row[1 + t] = float(x[t])          # 模板列 2..145 = t+1
        row[145] = float(np.sum(x))           # 全天购电量（列146）
        row[146] = float(l["cost_plan"])      # 全天购电费（列147）
        ws1.append(row)

    # ---- Sheet 2 充放电量 ----
    ws2 = out.create_sheet("充放电量")
    ws2.append(chg_header)
    for l in report_logs:
        d = parse_date(l["date"])
        c = l["plan_c"]                       # (144,) 计划充电
        r = l["plan_r"]                       # (144,) 计划放电
        s = l["plan_s"]                       # (145,) 计划 SOC
        windows = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
        labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
        for k, (a, b) in enumerate(windows):
            chg = float(np.sum(c[a:b]))
            dis = float(np.sum(r[a:b]))
            row = [None] * 6
            if k == 0:
                row[0] = d
                row[4] = dtime(0, 0)
                row[5] = float(s[0])
            elif k == 1:
                row[4] = "24:00"
                row[5] = float(s[T])
            row[1] = labels[k]
            row[2] = chg
            row[3] = dis
            ws2.append(row)

    # ---- Sheet 3 紧急购电量（连续时段合并，总纲 4.8） ----
    ws3 = out.create_sheet("紧急购电量")
    ws3.append(emg_header)
    for l in report_logs:
        d = parse_date(l["date"])
        e = l["settle_e"]
        t = 1
        while t <= T:
            if e[t - 1] > 1e-9:
                start = t
                while t <= T and e[t - 1] > 1e-9:
                    t += 1
                end = t - 1
                start_label = period_start[start]
                end_label = period_end[end]
                amount = float(np.sum(e[start - 1:end]))
                ws3.append([d, f"{start_label}-{end_label}", amount])
            else:
                t += 1

    out.save(OUT_XLSX)
    print(f"    result2.xlsx 已导出: {OUT_XLSX}")
    return OUT_XLSX


def verify_export(res):
    """聚合后与逐时段求和一致性复核。"""
    report_logs = res["report_logs"]
    total_plan = float(np.sum([np.sum(l["plan_x"]) for l in report_logs]))
    total_c = float(np.sum([np.sum(l["plan_c"]) for l in report_logs]))
    total_r = float(np.sum([np.sum(l["plan_r"]) for l in report_logs]))
    wb = openpyxl.load_workbook(OUT_XLSX)
    ws1 = wb["计划购电量"]
    col146_sum = float(np.sum([ws1.cell(r, 146).value or 0 for r in range(2, ws1.max_row + 1)]))
    ws2 = wb["充放电量"]
    chg_sum = float(np.sum([ws2.cell(r, 3).value or 0 for r in range(2, ws2.max_row + 1)]))
    dis_sum = float(np.sum([ws2.cell(r, 4).value or 0 for r in range(2, ws2.max_row + 1)]))
    ok = (abs(col146_sum - total_plan) < 1e-6 and abs(chg_sum - total_c) < 1e-6 and abs(dis_sum - total_r) < 1e-6)
    print(f"    复核：全天购电量列合计={col146_sum:.2f} vs 逐时段={total_plan:.2f}；"
          f"充电聚合={chg_sum:.2f} vs {total_c:.2f}；放电聚合={dis_sum:.2f} vs {total_r:.2f}；一致={ok}")
    return ok


def run_export():
    res, ds, tm = load_inputs()
    make_interval_mapping(ds, tm)
    build_workbook(res, ds, tm)
    verify_export(res)


def main() -> int:
    try:
        run_export()
        return 0
    except Exception as exc:
        print("=" * 60)
        print("P4 export 执行失败：")
        traceback.print_exc()
        print(f"错误信息: {exc}")
        return 1


if __name__ == "__main__":
    import sys
    sys.exit(main())
