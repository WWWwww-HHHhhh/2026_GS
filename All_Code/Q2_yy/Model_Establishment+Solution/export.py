# -*- coding: utf-8 -*-
"""
按官方模板导出 result2.xlsx（2026 数模 C 题 问题二）。

导出结构与官方模板保持一致：sheet 名、列名、列顺序与日期格式均按模板填写，
不新增 sheet 或列；相邻的十分钟紧急购电时段合并为连续区间后填写。
"""
import os
import pickle
import traceback
from datetime import datetime, time as dtime

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

from paths import Q2_ROOT, Q2_DATA_PROCESSING, Q2_TABLES, REPO_ROOT, RESULT2_TEMPLATE
Q2CODE = str(Q2_ROOT)
DATA_PROC = str(Q2_DATA_PROCESSING)
TABLES = str(Q2_TABLES)
DATA_TRANS = str(REPO_ROOT / "All_Code" / "Data_preprocessing" / "Data_transformation")
TEMPLATE_XLSX = str(RESULT2_TEMPLATE)
OUT_XLSX = os.path.join(TABLES, "result2.xlsx")
SPECIFIED_XLSX = os.path.join(TABLES, "specified_dates_results.xlsx")

T = 144


def load_inputs():
    with open(os.path.join(DATA_PROC, "q2_rolling_results.pkl"), "rb") as fh:
        res = pickle.load(fh)
    with open(os.path.join(DATA_PROC, "q2_dataset.pkl"), "rb") as fh:
        ds = pickle.load(fh)
    tm = pd.read_csv(os.path.join(DATA_TRANS, "time_map.csv"), encoding="utf-8-sig")
    return res, ds, tm


def make_interval_mapping(ds, tm):
    """内部时段序号 <-> 附件时间标签 <-> 模板列号 三列映射（真实时间口径）。
    规则：模板列 2..144 <- 本日 t=2..144；模板列 145（标签 0:00-0:10+1）<- 次日 t=1。"""
    rows = []
    tm_labels = {int(r.time_idx): str(r.template_col_label) for r in tm.itertuples()}
    for t in range(1, T + 1):
        row = tm[tm["time_idx"] == t].iloc[0]
        col = 145 if t == 1 else t
        label = tm_labels[144] if t == 1 else tm_labels[t - 1]
        rows.append([t, row["data_label"], label, col])
    df = pd.DataFrame(rows, columns=["内部时段序号", "附件时间标签(结束时刻)", "模板列标签", "模板列号"])
    path = os.path.join(TABLES, "interval_template_mapping.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="mapping")
    print(f"    时段-模板映射已写入: {path}")
    return df


def parse_date(ds_str):
    return datetime.strptime(ds_str, "%Y-%m-%d")


def build_workbook(res, ds, tm):
    """按官方模板结构填充三个 sheet。"""
    report_logs = res["report_logs"]
    # 官方模板结构
    tpl = openpyxl.load_workbook(TEMPLATE_XLSX)
    plan_tpl = tpl["计划购电量"]
    chg_tpl = tpl["充放电量"]
    emg_tpl = tpl["紧急购电量"]
    plan_header = [plan_tpl.cell(1, c).value for c in range(1, plan_tpl.max_column + 1)]
    chg_header = [chg_tpl.cell(1, c).value for c in range(1, chg_tpl.max_column + 1)]
    emg_header = [emg_tpl.cell(1, c).value for c in range(1, emg_tpl.max_column + 1)]
    if len(plan_header) != 147:
        raise ValueError(f"计划购电量模板列数应为 147，实际 {len(plan_header)}")

    out = openpyxl.Workbook()
    out.remove(out.active)

    # ---- Sheet 1 计划购电量 ----
    ws1 = out.create_sheet("计划购电量")
    ws1.append(plan_header)
    period_start = {int(r.time_idx): str(r.period_start) for r in tm.itertuples()}
    period_end = {int(r.time_idx): str(r.period_end) for r in tm.itertuples()}

    for i, l in enumerate(report_logs):
        d = parse_date(l["date"])
        x = l["plan_x"]                       # (144,)
        row = [d] + [None] * 146
        for t in range(1, T):                 # 本日 t=2..144 -> 列 2..144
            row[t] = float(x[t])
        nxt = report_logs[i + 1] if i + 1 < len(report_logs) else None
        if nxt is not None:
            if (parse_date(nxt["date"]) - d).days != 1:
                raise ValueError(f"report_logs 日期不连续: {l['date']} -> {nxt['date']}")
            row[144] = float(nxt["plan_x"][0])   # 末列（列145）<- 次日 t=1
        else:
            row[144] = None                     # 最后一行（2025-12-31）末列留空
        row[145] = float(np.sum(x))           # 全天购电量（列146）
        row[146] = float(l["cost_plan"])      # 全天购电费（列147）
        ws1.append(row)

    # ---- Sheet 2 充放电量 ----
    ws2 = out.create_sheet("充放电量")
    ws2.append(chg_header)
    for l in report_logs:
        d = parse_date(l["date"])
        # 与紧急购电使用同一次因果执行轨迹，保证提交表可由电量平衡复算。
        c = l["settle_c_actual"]              # (144,) 实际执行充电
        r = l["settle_r_actual"]              # (144,) 实际执行放电
        s = l["settle_s_actual"]              # (145,) 实际 SOC
        windows = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
        labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
        for k, (a, b) in enumerate(windows):
            chg = float(np.sum(c[a:b]))
            dis = float(np.sum(r[a:b]))
            row = [None] * 6
            if k == 0:
                row[0] = d
                row[4] = dtime(0, 0)
                row[5] = float(s[0])
            elif k == 1:
                row[4] = "24:00"
                row[5] = float(s[T])
            row[1] = labels[k]
            row[2] = chg
            row[3] = dis
            ws2.append(row)

    # ---- Sheet 3 紧急购电量（相邻时段合并为连续区间） ----
    ws3 = out.create_sheet("紧急购电量")
    ws3.append(emg_header)
    for l in report_logs:
        d = parse_date(l["date"])
        e = l["settle_e"]
        t = 1
        while t <= T:
            if e[t - 1] > 1e-9:
                start = t
                while t <= T and e[t - 1] > 1e-9:
                    t += 1
                end = t - 1
                start_label = period_start[start]
                end_label = period_end[end]
                amount = float(np.sum(e[start - 1:end]))
                ws3.append([d, f"{start_label}-{end_label}", amount])
            else:
                t += 1

    out.save(OUT_XLSX)
    print(f"    result2.xlsx 已导出: {OUT_XLSX}")
    return OUT_XLSX


def verify_export(res):
    """聚合后与逐时段求和一致性复核；并检查 2025-12-31 行末列按要求留空。"""
    report_logs = res["report_logs"]
    total_plan = float(np.sum([np.sum(l["plan_x"]) for l in report_logs]))
    total_c = float(np.sum([np.sum(l["settle_c_actual"]) for l in report_logs]))
    total_r = float(np.sum([np.sum(l["settle_r_actual"]) for l in report_logs]))
    total_e = float(np.sum([np.sum(l["settle_e"]) for l in report_logs]))
    wb = openpyxl.load_workbook(OUT_XLSX)
    ws1 = wb["计划购电量"]
    col146_sum = float(np.sum([ws1.cell(r, 146).value or 0 for r in range(2, ws1.max_row + 1)]))
    last_row_c145 = ws1.cell(ws1.max_row, 145).value
    ws2 = wb["充放电量"]
    chg_sum = float(np.sum([ws2.cell(r, 3).value or 0 for r in range(2, ws2.max_row + 1)]))
    dis_sum = float(np.sum([ws2.cell(r, 4).value or 0 for r in range(2, ws2.max_row + 1)]))
    ws3 = wb["紧急购电量"]
    emg_sum = float(np.sum([ws3.cell(r, 3).value or 0 for r in range(2, ws3.max_row + 1)]))
    ok = (abs(col146_sum - total_plan) < 1e-6 and abs(chg_sum - total_c) < 1e-6
          and abs(dis_sum - total_r) < 1e-6 and abs(emg_sum - total_e) < 1e-6)
    print(f"    复核：全天购电量列合计={col146_sum:.2f} vs 逐时段={total_plan:.2f}；"
          f"充电聚合={chg_sum:.2f} vs {total_c:.2f}；放电聚合={dis_sum:.2f} vs {total_r:.2f}；"
          f"紧急购电={emg_sum:.2f} vs {total_e:.2f}；一致={ok}")
    print(f"    复核：末行（{ws1.cell(ws1.max_row, 1).value.strftime('%Y-%m-%d')}）末列值={last_row_c145}（应为空 None）")
    if not ok or last_row_c145 is not None:
        raise RuntimeError("result2.xlsx 独立聚合复核失败")
    return True


def export_specified_dates(res, tm):
    """按题目表1、表2、表3口径汇总四个指定日期，供论文直接引用。"""
    wanted = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
    by_date = {l["date"]: l for l in res["report_logs"]}
    if any(d not in by_date for d in wanted):
        raise ValueError("报告结果缺少题目指定日期")
    point_labels = ["10:00-10:10", "12:00-12:10", "14:00-14:10",
                    "16:00-16:10", "18:00-18:10", "20:00-20:10"]
    point_index = [60, 72, 84, 96, 108, 120]
    plan_rows, storage_rows, emergency_rows = [], [], []
    starts = {int(r.time_idx): str(r.period_start) for r in tm.itertuples()}
    ends = {int(r.time_idx): str(r.period_end) for r in tm.itertuples()}
    windows = [(0, 24, "0:00-4:00"), (24, 48, "4:00-8:00"),
               (48, 72, "8:00-12:00"), (72, 96, "12:00-16:00"),
               (96, 120, "16:00-20:00"), (120, 144, "20:00-24:00")]
    for d in wanted:
        l = by_date[d]
        row = {"日期": d}
        for label, idx in zip(point_labels, point_index):
            row[label + "购电量(kWh)"] = float(l["plan_x"][idx])
        row["全天计划购电量(kWh)"] = float(np.sum(l["plan_x"]))
        row["全天计划购电费(元)"] = float(l["cost_plan"])
        row["全天紧急购电费(元)"] = float(l["cost_emergency"])
        row["全天总费用(元)"] = float(l["cost_total"])
        plan_rows.append(row)
        for a, b, label in windows:
            storage_rows.append({"日期": d, "时间段": label,
                                 "实际充电量(kWh)": float(np.sum(l["settle_c_actual"][a:b])),
                                 "实际放电量(kWh)": float(np.sum(l["settle_r_actual"][a:b])),
                                 "0:00储电量(kWh)": float(l["settle_s_actual"][0]),
                                 "24:00储电量(kWh)": float(l["settle_s_actual"][-1])})
        e = l["settle_e"]
        t = 1
        while t <= T:
            if e[t-1] <= 1e-9:
                t += 1; continue
            a = t
            while t <= T and e[t-1] > 1e-9:
                t += 1
            b = t-1
            emergency_rows.append({"日期": d, "紧急购电时间段": f"{starts[a]}-{ends[b]}",
                                   "紧急购电量(kWh)": float(np.sum(e[a-1:b]))})
    with pd.ExcelWriter(SPECIFIED_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(plan_rows).to_excel(writer, index=False, sheet_name="表1购电")
        pd.DataFrame(storage_rows).to_excel(writer, index=False, sheet_name="表2储能")
        pd.DataFrame(emergency_rows, columns=["日期", "紧急购电时间段", "紧急购电量(kWh)"]).to_excel(
            writer, index=False, sheet_name="表3紧急购电")
    print(f"    指定日期论文表已写入: {SPECIFIED_XLSX}")


def run_export():
    res, ds, tm = load_inputs()
    make_interval_mapping(ds, tm)
    build_workbook(res, ds, tm)
    verify_export(res)
    export_specified_dates(res, tm)


def main() -> int:
    try:
        run_export()
        return 0
    except Exception as exc:
        print("=" * 60)
        print("result2 导出失败：")
        traceback.print_exc()
        print(f"错误信息: {exc}")
        return 1


if __name__ == "__main__":
    import sys
    sys.exit(main())
