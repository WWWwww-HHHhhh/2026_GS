# -*- coding: utf-8 -*-
"""
问题1：含弃光和储能边际价值的连续 LP（HiGHS 求解）

输入数据：Data/附件/附件1.xlsx（144 个 10 min 时段的电价、小区负载、光伏预测功率）
         Data/附件/附件5/result1.xlsx（官方结果模板）
运行方法：python q1_model.py（依赖 numpy/pandas/scipy/openpyxl；Python 3.13 验证通过）
输出位置：All_Code/Q1/Tables/（result1.xlsx、q1_timeseries.csv、q1_summary.json）

模型（确定性单日，功率 × 1/6 转为时段电量 kWh）：
    min  J1 = sum_t pi_t * x_t + eps * sum_t (c_t + r_t)
    s.t. x_t + (G_t - w_t) + r_t = L_t + c_t        电量平衡（含弃光）
         s_t = s_{t-1} + eta_c*c_t - r_t/eta_r     储能递推
         S_min <= s_t <= S_max, s_0 = s_T = 6000   SOC 边界（问题1硬周期）
         0 <= c_t, r_t <= 833.3333                  5000 kW 折算时段电量
         0 <= w_t <= G_t, x_t >= 0
真实购电费用 C1 = sum_t pi_t * x_t（不含 eps 惩罚），单独报告。
对偶：mu_bal = 时段边际供电成本；储能边际价值 = -mu_soc。符号约定经数值扰动核验。
"""

import json
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linprog

# ---------------------------------------------------------------- 路径与常量
Q1_DIR = Path(__file__).resolve().parents[1]           # All_Code/Q1
REPO_ROOT = Path(__file__).resolve().parents[3]        # 仓库根目录
DATA_XLSX = REPO_ROOT / "Data" / "附件" / "附件1.xlsx"
TEMPLATE_XLSX = REPO_ROOT / "Data" / "附件" / "附件5" / "result1.xlsx"
RESULTS = Q1_DIR / "Tables"
RESULTS.mkdir(exist_ok=True)

T = 144                      # 每日 10 min 时段数
DT = 1.0 / 6.0               # 时段长度 h
S_MIN, S_MAX = 1200.0, 10800.0
S0 = 6000.0                  # 0:00 与 24:00 储电量（问题1硬边界）
P_MAX_KWH = 5000.0 * DT      # 单时段最大充/放电量 833.3333 kWh
ETA_C = ETA_R = 0.9
EPS = 1e-3                   # 储能吞吐微小惩罚，元/kWh（约为最低电价的 0.27%）


def load_data():
    """读取附件1：144 个时段的电价(元/kWh)、负载(kW)、光伏预测(kW)，功率乘 1/6 转电量。"""
    df = pd.read_excel(DATA_XLSX, header=0)
    df.columns = ["time_label", "price", "load_kw", "pv_kw"]
    assert len(df) == T, f"附件1应有 {T} 行，实际 {len(df)} 行"
    L = df["load_kw"].to_numpy(float) * DT   # kWh/时段
    G = df["pv_kw"].to_numpy(float) * DT
    pi = df["price"].to_numpy(float)
    return df, pi, L, G


def build_and_solve(pi, L, G, eps=EPS):
    """构建并求解主 LP。变量排序：x(0:T), c(T:2T), r(2T:3T), w(3T:4T), s(4T:5T)。"""
    n = 5 * T
    c_obj = np.concatenate([pi, np.full(T, eps), np.full(T, eps), np.zeros(2 * T)])

    A_eq, b_eq = [], []

    # 电量平衡：x_t - w_t + r_t - c_t = L_t - G_t
    for t in range(T):
        row = np.zeros(n)
        row[t] = 1.0            # x_t
        row[T + t] = -1.0       # -c_t
        row[2 * T + t] = 1.0    # +r_t
        row[3 * T + t] = -1.0   # -w_t
        A_eq.append(row)
        b_eq.append(L[t] - G[t])

    # SOC 递推：s_t - eta_c*c_t + r_t/eta_r - s_{t-1} = 0（s_0 = S0 为常数移项）
    for t in range(T):
        row = np.zeros(n)
        row[4 * T + t] = 1.0            # s_t
        row[T + t] = -ETA_C             # -eta_c*c_t
        row[2 * T + t] = 1.0 / ETA_R    # +r_t/eta_r
        if t > 0:
            row[4 * T + t - 1] = -1.0   # -s_{t-1}
        A_eq.append(row)
        b_eq.append(S0 if t == 0 else 0.0)

    # 硬周期边界：s_T = S0
    row = np.zeros(n)
    row[4 * T + T - 1] = 1.0
    A_eq.append(row)
    b_eq.append(S0)

    bounds = (
        [(0.0, None)] * T +                       # x
        [(0.0, P_MAX_KWH)] * T +                  # c
        [(0.0, P_MAX_KWH)] * T +                  # r
        [(0.0, float(G[t])) for t in range(T)] +  # w <= G_t
        [(S_MIN, S_MAX)] * T                      # s
    )

    res = linprog(c_obj, A_eq=np.array(A_eq), b_eq=np.array(b_eq),
                  bounds=bounds, method="highs")
    if res.status != 0:
        raise RuntimeError(f"LP 求解失败: status={res.status}, message={res.message}")

    sol = {
        "x": res.x[0:T], "c": res.x[T:2 * T], "r": res.x[2 * T:3 * T],
        "w": res.x[3 * T:4 * T], "s": res.x[4 * T:5 * T],
        "J1": float(res.fun),
        "C1": float(pi @ res.x[0:T]),
        "mu_bal": np.array(res.eqlin.marginals[0:T]),
        "mu_soc": np.array(res.eqlin.marginals[T:2 * T]),
    }
    return sol


def perturbation_check(pi, L, G, sol):
    """数值扰动核验对偶符号约定：
    1) 某时段负荷 +1 kWh，dC1 应等于该时段平衡对偶（即边际供电成本）；
    2) 周期边界 S0 + 1 kWh，观察 SOC 边界价值量级。"""
    t_probe = 100  # 代表时段
    L2 = L.copy(); L2[t_probe] += 1.0
    sol2 = build_and_solve(pi, L2, G)
    grad_bal = sol2["C1"] - sol["C1"]

    global S0
    S0_orig = S0
    S0 = S0_orig + 1.0
    sol3 = build_and_solve(pi, L, G)
    S0 = S0_orig
    grad_soc = sol3["C1"] - sol["C1"]
    return {"t_probe": t_probe, "grad_bal": grad_bal, "grad_soc": grad_soc}


def validate(sol, L, G):
    """校验层一：平衡残差、SOC 递推残差、边界检查。"""
    x, c, r, w, s = sol["x"], sol["c"], sol["r"], sol["w"], sol["s"]
    bal = x + (G - w) + r - L - c
    s_prev = np.concatenate([[S0], s[:-1]])
    soc_res = s - s_prev - ETA_C * c + r / ETA_R
    checks = {
        "max_abs_balance_residual_kwh": float(np.abs(bal).max()),
        "max_abs_soc_recursion_residual_kwh": float(np.abs(soc_res).max()),
        "soc_min": float(s.min()), "soc_max": float(s.max()),
        "soc_within_bounds": bool(s.min() >= S_MIN - 1e-6 and s.max() <= S_MAX + 1e-6),
        "s0_equals_sT": bool(abs(s[-1] - S0) < 1e-6),
        "max_charge_kwh": float(c.max()), "max_discharge_kwh": float(r.max()),
        "charge_discharge_within_power_limit": bool(
            c.max() <= P_MAX_KWH + 1e-6 and r.max() <= P_MAX_KWH + 1e-6),
        "curtailment_nonnegative_and_le_pv": bool(
            w.min() >= -1e-9 and np.all(w <= G + 1e-6)),
        "simultaneous_charge_discharge_max_product": float((c * r).max()),
    }
    return checks


def export_result1(sol):
    """按模板逐行映射填充 result1.xlsx（不修改模板表头）。
    映射：内部时段序号 t（1..144，与附件1第 t 行同序）-> 模板“计划购电量”第 t 行。
    专项核对：模板 144 行全部填满；填充总量与逐时段求和一致。"""
    out = RESULTS / "result1.xlsx"
    shutil.copy(TEMPLATE_XLSX, out)

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb["计划购电量"]
    assert ws.max_row - 1 == T, f"模板应有 {T} 行数据，实际 {ws.max_row - 1}"
    for t in range(T):
        ws.cell(row=t + 2, column=2, value=float(sol["x"][t]))
    filled = sum(ws.cell(row=t + 2, column=2).value for t in range(T))
    assert abs(filled - sol["x"].sum()) < 1e-6, "模板填充总量与逐时段求和不一致"

    ws2 = wb["充放电量"]
    for blk in range(6):  # 六个 4 h 区间 = 24 个时段
        lo, hi = blk * 24, (blk + 1) * 24
        ws2.cell(row=blk + 2, column=2, value=float(sol["c"][lo:hi].sum()))
        ws2.cell(row=blk + 2, column=3, value=float(sol["r"][lo:hi].sum()))
    ws2.cell(row=2, column=5, value=S0)                 # 0:00 储电量
    ws2.cell(row=3, column=5, value=float(sol["s"][-1]))  # 24:00 储电量
    wb.save(out)
    return out


def main():
    df, pi, L, G = load_data()
    print(f"数据: T={T}, 电价[{pi.min():.4f},{pi.max():.4f}] 元/kWh, "
          f"负荷日均功率 {L.sum()/24:.1f} kW, 光伏峰值 {G.max()/DT:.1f} kW, "
          f"光伏日发电量 {G.sum():.1f} kWh, 负荷日用电量 {L.sum():.1f} kWh")

    sol = build_and_solve(pi, L, G)
    checks = validate(sol, L, G)
    pert = perturbation_check(pi, L, G, sol)
    out = export_result1(sol)

    # 逐时段结果存档（论文数值可溯源）
    detail = pd.DataFrame({
        "time_label": df["time_label"], "price": pi,
        "L_kwh": L, "G_kwh": G,
        "x_plan_kwh": sol["x"], "c_charge_kwh": sol["c"],
        "r_discharge_kwh": sol["r"], "w_curtail_kwh": sol["w"],
        "soc_kwh": sol["s"], "mu_bal": sol["mu_bal"], "mu_soc": sol["mu_soc"],
    })
    detail.to_csv(RESULTS / "q1_timeseries.csv", index=False, encoding="utf-8-sig")

    summary = {
        "C1_true_purchase_cost_yuan": sol["C1"],
        "J1_objective_with_eps_yuan": sol["J1"],
        "total_purchase_kwh": float(sol["x"].sum()),
        "total_curtailment_kwh": float(sol["w"].sum()),
        "storage_throughput_kwh": float(sol["c"].sum() + sol["r"].sum()),
        "charge_kwh": float(sol["c"].sum()), "discharge_kwh": float(sol["r"].sum()),
        "eps_yuan_per_kwh": EPS,
        "validation": checks,
        "perturbation_check": pert,
        "mu_soc_stats": {"min": float(sol["mu_soc"].min()),
                          "max": float(sol["mu_soc"].max())},
        "mu_bal_stats": {"min": float(sol["mu_bal"].min()),
                          "max": float(sol["mu_bal"].max())},
    }
    with open(RESULTS / "q1_summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\nC1 真实购电费用 = {sol['C1']:.2f} 元;  全天购电量 = {sol['x'].sum():.2f} kWh")
    print(f"弃光 = {sol['w'].sum():.2f} kWh;  储能吞吐 = {summary['storage_throughput_kwh']:.2f} kWh")
    print(f"最大平衡残差 = {checks['max_abs_balance_residual_kwh']:.2e} kWh; "
          f"SOC ∈ [{checks['soc_min']:.1f}, {checks['soc_max']:.1f}]")
    print(f"同时充放最大乘积 = {checks['simultaneous_charge_discharge_max_product']:.2e}")
    print(f"扰动核验: t={pert['t_probe']} 负荷边际成本 {pert['grad_bal']:.4f} 元/kWh "
          f"(pi={pi[pert['t_probe']]:.4f}); SOC 边界价值实测 {pert['grad_soc']:.4f} 元/kWh")
    print(f"已输出: {out}")


if __name__ == "__main__":
    sys.exit(main())
