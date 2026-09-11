# -*- coding: utf-8 -*-
"""优化版 Q2 自动校验（真实数据核对，不含任何伪造结论）。"""
import os, pickle
from datetime import datetime
import numpy as np
import openpyxl
from paths import Q2_DATA_PROCESSING, Q2_TABLES

T = 144
TOL = 1e-5
SMIN, SMAX, CAP = 1200.0, 10800.0, 833.3333333333334

def load():
    with open(os.path.join(Q2_DATA_PROCESSING, "net_rolling_results.pkl"), "rb") as fh:
        res = pickle.load(fh)
    with open(os.path.join(Q2_DATA_PROCESSING, "q2_dataset.pkl"), "rb") as fh:
        ds = pickle.load(fh)
    return res, ds

def run():
    res, ds = load()
    logs = res["report_logs"]; full = res["full_logs"]
    checks = []
    # 1 结算平衡
    mx = max(float(np.max(np.abs(l["settle_y"] + l["settle_e"] + l["settle_g"] + l["settle_r_actual"] - l["load_actual"] - l["settle_c_actual"]))) for l in logs)
    checks.append(("结算电量平衡残差<=1e-5", mx <= TOL, f"max={mx:.3e}"))
    # 2 SOC
    bad = sum(1 for l in logs if np.any((l["plan_s"] < SMIN - 1e-6) | (l["plan_s"] > SMAX + 1e-6)) or np.any((l["settle_s_actual"] < SMIN - 1e-6) | (l["settle_s_actual"] > SMAX + 1e-6)))
    checks.append(("SOC∈[1200,10800]", bad == 0, f"越界天数={bad}"))
    # 3 充放电上限
    bad3 = sum(1 for l in logs if np.any(l["plan_c"] > CAP + 1e-6) or np.any(l["plan_r"] > CAP + 1e-6) or np.any(l["settle_c_actual"] > CAP + 1e-6) or np.any(l["settle_r_actual"] > CAP + 1e-6))
    checks.append(("充放电<=833.333", bad3 == 0, f"超限天数={bad3}"))
    # 4 跨日 SOC
    mism = sum(1 for k in range(len(full) - 1) if abs(full[k]["final_soc"] - full[k + 1]["s0"]) > 1e-6)
    checks.append(("跨日SOC传递一致", mism == 0, f"不一致天数={mism}"))
    # 5 因果性：残差块数等于决策日前一天可用天数
    ok5 = all(l["residual_blocks"] == l["day_index"] - 1 for l in logs)
    checks.append(("因果性：残差库仅含历史", ok5, "全部满足" if ok5 else "存在异常"))
    # 6 弃光公式
    ok6 = all(np.max(np.abs(l["settle_w"] - (l["pv_actual"] - l["settle_g"]))) < 1e-6 for l in logs)
    checks.append(("弃光=光伏可用-消纳", ok6, "全部满足" if ok6 else "存在异常"))
    # 7 紧急购电5倍复算
    ok7 = all(abs(l["cost_emergency"] - float(np.sum(5.0 * l["price"] * l["settle_e"]))) < 1e-6 for l in logs)
    checks.append(("紧急购电5倍电价复算", ok7, "全部满足" if ok7 else "存在异常"))
    # 8 报告费用口径
    ok8 = all(abs(l["cost_total"] - (l["cost_plan"] + l["cost_emergency"])) < 1e-6 for l in logs)
    checks.append(("报告费用=计划+紧急", ok8, "全部满足" if ok8 else "存在异常"))
    # 9 样本量
    checks.append(("报告期样本量=334", len(logs) == 334, f"n={len(logs)}"))
    # 10 result2 总量一致
    wb = openpyxl.load_workbook(os.path.join(Q2_TABLES, "result2.xlsx"))
    ws = wb["计划购电量"]
    colsum = float(np.sum([ws.cell(r, 146).value or 0 for r in range(2, ws.max_row + 1)]))
    plansum = float(np.sum([np.sum(l["plan_x"]) for l in logs]))
    checks.append(("result2计划购电总量一致", abs(colsum - plansum) < 1e-6, f"文件={colsum:.2f} vs 日志={plansum:.2f}"))

    lines = []
    all_pass = True
    for name, ok, ev in checks:
        all_pass = all_pass and ok
        lines.append(f"[{'PASS' if ok else 'FAIL'}] {name}: {ev}")
    with open(os.path.join(Q2_TABLES, "validation_report.txt"), "w", encoding="utf-8") as fh:
        fh.write(f"问题二 优化版自动校验报告（生成于 {datetime.now():%Y-%m-%d %H:%M:%S}）\n")
        fh.write(f"总体结果: {'全部 PASS' if all_pass else '存在 FAIL'}\n")
        fh.write("-" * 60 + "\n" + "\n".join(lines) + "\n")
    print("总体:", "ALL PASS" if all_pass else "HAS FAIL")
    for ln in lines:
        print(ln)
    return all_pass

if __name__ == "__main__":
    run()
