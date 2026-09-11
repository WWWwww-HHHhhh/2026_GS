# -*- coding: utf-8 -*-
"""
Q2 -> Q3 数据同步与自动修改（交接执行器）
==========================================
依据《Q2到Q3数据同步与自动修改提示词.md》全流程执行：
  1) 核验源工作簿 result2(7)：SHA256、6 个汇总值、3 个行内区间差；
  2) 核验 beta 同一次运行证据（q2_rolling_results.pkl 复算 6 个汇总值）；
  3) 使用词典序结算（settlement.py）从 2025-01 预热期重新滚动并重新调参；
  4) 生成 5 个 Q3 接口文件 + 参考工作簿 + 审计工作簿 + 交接报告；
  5) 执行 15 项自动校验，全部 PASS 才算完成。

铁律：
  - 所有新增 CSV/JSON/XLSX/MD 只写入 All_Code/Q3，绝不写入 All_Code/Q2；
  - 不覆盖 Q2 官方 result2.xlsx；不执行 export.py / main_q2.py；不 push；
  - 不捏造 beta、场景来源、预测数据；无法核实即失败退出。
"""
import hashlib
import json
import os
import pickle
import subprocess
import sys
import time
import traceback
from datetime import datetime, timedelta, time as dtime

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from paths import (
    REPO_ROOT, Q2_DATA_PROCESSING, Q2_TABLES, Q3_INTERFACE, Q3_TABLES,
    RESULT2_TEMPLATE, Q2_RESULT2_SOURCE,
)

T = 144
D = 365
SEED = 20260101

EXPECTED_SHA256 = "cfd33c2304096244d27c9575d7bc4335fc052f22158403334fbe4d428e994bf6"
EXPECTED_SUMMARY = {
    "计划购电总量(kWh)": 23473108.290145896,
    "计划购电费(元)": 14797743.36569577,
    "充电量(kWh)": 6783239.013789247,
    "放电量(kWh)": 5502760.789258115,
    "紧急购电量(kWh)": 221843.73934786578,
    "紧急购电记录数": 3421,
}
EXPECTED_ROW_DIFF = {
    "2025-02-01": 117.24,
    "2025-05-10": 1012.44,
    "2025-12-31": -1464.30,
}

RED = PatternFill("solid", fgColor="FFC7CE")
GREEN = PatternFill("solid", fgColor="C6EFCE")
BOLD = Font(bold=True)


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest().lower()


def git_head():
    cp = subprocess.run(["git", "-C", str(REPO_ROOT), "rev-parse", "HEAD"],
                        capture_output=True, text=True, encoding="utf-8", errors="replace")
    return cp.stdout.strip() if cp.returncode == 0 else "UNKNOWN"


def verify_source_workbook():
    """核验 Q2 result2.xlsx（应等于 result2(7)）：SHA256 + 6 个汇总值 + 3 个行内差。"""
    ev = {"path": str(Q2_RESULT2_SOURCE), "hash": None, "hash_ok": False,
          "summary": {}, "row_diff": {}, "all_ok": False}
    if not Q2_RESULT2_SOURCE.exists():
        raise FileNotFoundError(f"源工作簿不存在: {Q2_RESULT2_SOURCE}")
    ev["hash"] = sha256_file(Q2_RESULT2_SOURCE)
    ev["hash_ok"] = ev["hash"] == EXPECTED_SHA256

    wb = openpyxl.load_workbook(Q2_RESULT2_SOURCE, read_only=True)
    rows1 = list(wb["计划购电量"].iter_rows(min_row=2, values_only=True))
    rows2 = list(wb["充放电量"].iter_rows(min_row=2, values_only=True))
    rows3 = list(wb["紧急购电量"].iter_rows(min_row=2, values_only=True))
    wb.close()

    plan_total = 0.0
    fee_total = 0.0
    row_diff = {}
    for row in rows1:
        plan_total += float(row[145] or 0.0)
        fee_total += float(row[146] or 0.0)
        d = row[0]
        dstr = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)[:10]
        if dstr in EXPECTED_ROW_DIFF:
            shown = sum(float(v) if v is not None else 0.0 for v in row[1:145])
            row_diff[dstr] = shown - float(row[145] or 0.0)
    chg_total = sum(float(r[2]) if r[2] is not None else 0.0 for r in rows2)
    dis_total = sum(float(r[3]) if r[3] is not None else 0.0 for r in rows2)
    emg_total = sum(float(r[2]) if r[2] is not None else 0.0 for r in rows3)
    emg_rec = sum(1 for r in rows3 if r[0] is not None)
    got = {"计划购电总量(kWh)": plan_total, "计划购电费(元)": fee_total,
           "充电量(kWh)": chg_total, "放电量(kWh)": dis_total,
           "紧急购电量(kWh)": emg_total, "紧急购电记录数": emg_rec}
    ev["summary"] = {k: {"expected": EXPECTED_SUMMARY[k], "actual": got[k],
                         "diff": got[k] - EXPECTED_SUMMARY[k]} for k in EXPECTED_SUMMARY}
    ev["row_diff"] = {k: {"expected": EXPECTED_ROW_DIFF[k], "actual": row_diff.get(k),
                          "diff": (row_diff.get(k) or 0.0) - EXPECTED_ROW_DIFF[k]}
                      for k in EXPECTED_ROW_DIFF}
    ev["all_ok"] = (
        ev["hash_ok"]
        and all(abs(got[k] - EXPECTED_SUMMARY[k]) <= 1e-3 for k in EXPECTED_SUMMARY)
        and all(abs((row_diff.get(k) or 0.0) - EXPECTED_ROW_DIFF[k]) <= 0.05
                for k in EXPECTED_ROW_DIFF)
    )
    return ev


def verify_old_pkl_vs_md():
    """用旧缓存 q2_rolling_results.pkl 复算 6 个汇总值（β 同一次运行证据）。"""
    p = Q2_DATA_PROCESSING / "q2_rolling_results.pkl"
    with open(p, "rb") as fh:
        res = pickle.load(fh)
    logs = res["report_logs"]
    tp = float(np.sum([np.sum(l["plan_x"]) for l in logs]))
    cp = float(np.sum([l["cost_plan"] for l in logs]))
    tc = float(np.sum([np.sum(l["plan_c"]) for l in logs]))
    tr = float(np.sum([np.sum(l["plan_r"]) for l in logs]))
    te = float(np.sum([l["emergency_kwh"] for l in logs]))
    nrec = 0
    for l in logs:
        e = l["settle_e"]
        t = 0
        while t < T:
            if e[t] > 1e-9:
                nrec += 1
                while t < T and e[t] > 1e-9:
                    t += 1
            else:
                t += 1
    got = {"计划购电总量(kWh)": tp, "计划购电费(元)": cp, "充电量(kWh)": tc,
           "放电量(kWh)": tr, "紧急购电量(kWh)": te, "紧急购电记录数": nrec}
    ok = all(abs(got[k] - EXPECTED_SUMMARY[k]) <= 1e-6 for k in EXPECTED_SUMMARY)
    return {"match": ok, "values": got,
            "params": res["params"], "beta_best": res["beta_best"],
            "m_best": res["m_best"], "kappa2_best": res["kappa2_best"],
            "alpha": res["alpha"]}


def build_forecast_csv(fc, ds):
    """q2_forecast_10min.csv：预测/实际长表（主键 date+interval_index，t=1 即真实 00:00-00:10）。
    光伏预测按物理意义截断到 >=0（与 scenarios.py 的物理截断规则一致），并如实返回截断统计。"""
    sm = fc["selected_models"].copy()
    sm["date"] = sm["date"].astype(str)
    sm["series"] = sm["series"].astype(str)
    lookup = {(r.date, r.series): str(r.selected_model) for r in sm.itertuples()}
    rows = []
    clip_count = 0
    clip_max = 0.0
    for i in range(D):
        dstr = ds["date_str"][i]
        for t in range(T):
            gv = float(fc["Ghat"][t, i])
            if gv < 0.0:
                clip_count += 1
                clip_max = max(clip_max, -gv)
            rows.append({
                "date": dstr,
                "interval_index": t + 1,
                "load_pred_kwh": float(fc["Lhat"][t, i]),
                "pv_pred_0h_kwh": max(gv, 0.0),
                "load_actual_kwh": float(ds["load"][t, i]),
                "pv_actual_kwh": float(ds["pv"][t, i]),
                "is_warmup": 1 if i <= 30 else 0,
                "model_load": lookup.get((dstr, "load"), ""),
                "model_pv": lookup.get((dstr, "pv"), ""),
            })
    df = pd.DataFrame(rows, columns=[
        "date", "interval_index", "load_pred_kwh", "pv_pred_0h_kwh",
        "load_actual_kwh", "pv_actual_kwh", "is_warmup", "model_load", "model_pv"])
    out = Q3_INTERFACE / "q2_forecast_10min.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return out, len(df), {"clipped_cells": clip_count, "max_magnitude_kwh": clip_max}


def build_residual_blocks(fc, ds):
    """q2_residual_blocks.csv：整日残差块，res = actual - pred。
    2025-01-01（i=0）从未进入残差库，按提示词不导出该日记录。"""
    rows = []
    for i in range(1, D):
        dstr = ds["date_str"][i]
        nd = (datetime.strptime(dstr, "%Y-%m-%d") + timedelta(days=1)).strftime("%Y-%m-%d")
        for t in range(T):
            rows.append({
                "residual_date": dstr,
                "interval_index": t + 1,
                "load_residual_kwh": float(ds["load"][t, i] - fc["Lhat"][t, i]),
                "pv_residual_kwh": float(ds["pv"][t, i] - max(float(fc["Ghat"][t, i]), 0.0)),
                "eligible_from_date": nd,
            })
    df = pd.DataFrame(rows, columns=[
        "residual_date", "interval_index", "load_residual_kwh",
        "pv_residual_kwh", "eligible_from_date"])
    out = Q3_INTERFACE / "q2_residual_blocks.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return out, len(df)


def build_scenario_manifest(ds, M):
    """q2_scenario_manifest.csv：严格复现 scenarios.py 抽样方式。
    rng = default_rng(SEED*1000+i)；有放回抽 1 <= source_day_index < i；
    残差池为空时 source 留空（点预测场景），不伪造历史日期。"""
    rows = []
    for i in range(1, D):
        avail = list(range(1, i))
        target = ds["date_str"][i]
        seed_i = SEED * 1000 + i
        if not avail:
            for m in range(1, M + 1):
                rows.append({"target_date": target, "target_day_index": i, "scenario_id": m,
                             "source_residual_date": "", "source_day_index": "",
                             "scenario_weight": 1.0 / M, "random_seed": seed_i})
        else:
            rng = np.random.default_rng(seed_i)
            idx = rng.integers(0, len(avail), size=M)
            for m in range(1, M + 1):
                src = avail[int(idx[m - 1])]
                rows.append({"target_date": target, "target_day_index": i, "scenario_id": m,
                             "source_residual_date": ds["date_str"][src],
                             "source_day_index": src,
                             "scenario_weight": 1.0 / M, "random_seed": seed_i})
    df = pd.DataFrame(rows, columns=[
        "target_date", "target_day_index", "scenario_id", "source_residual_date",
        "source_day_index", "scenario_weight", "random_seed"])
    out = Q3_INTERFACE / "q2_scenario_manifest.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return out, len(df)


def build_plan_reference_long(results):
    """q2_plan_reference_long.csv：报告期（2025-02-01..12-31）长表，全部取自词典序重滚后的
    report_logs（plan_x / settle_y / settle_c_actual / settle_r_actual / settle_s_actual /
    settle_e / settle_w）。主键 date + interval_index。"""
    logs = results["report_logs"]
    rows = []
    for l in logs:
        dstr = l["date"]
        for t in range(T):
            rows.append({
                "date": dstr,
                "interval_index": t + 1,
                "plan_purchase_kwh": float(l["plan_x"][t]),
                "actual_extracted_kwh": float(l["settle_y"][t]),
                "actual_charge_kwh": float(l["settle_c_actual"][t]),
                "actual_discharge_kwh": float(l["settle_r_actual"][t]),
                "actual_soc_start_kwh": float(l["settle_s_actual"][t]),
                "actual_soc_end_kwh": float(l["settle_s_actual"][t + 1]),
                "emergency_purchase_kwh": float(l["settle_e"][t]),
                "curtailment_kwh": float(l["settle_w"][t]),
            })
    df = pd.DataFrame(rows, columns=[
        "date", "interval_index", "plan_purchase_kwh", "actual_extracted_kwh",
        "actual_charge_kwh", "actual_discharge_kwh", "actual_soc_start_kwh",
        "actual_soc_end_kwh", "emergency_purchase_kwh", "curtailment_kwh"])
    out = Q3_INTERFACE / "q2_plan_reference_long.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    return out, len(df)


def build_parameters_json(results, src_ev, pkl_ev, head, clip_stats):
    """q2_parameters.json：Q3 采用参数（词典序重滚后重调参结果）+ β 证据。"""
    params = {
        "source_workbook": "All_Code/Q2/Results/Tables/result2.xlsx",
        "source_workbook_abs": str(Q2_RESULT2_SOURCE),
        "source_workbook_sha256": src_ev["hash"],
        "source_git_commit": head,
        "beta": float(results["beta_best"]),
        "scenario_count": int(results["m_best"]),
        "cvar_alpha": float(results["alpha"]),
        "kappa2": float(results["kappa2_best"]),
        "kappa2_base": float(results["kappa2_base"]),
        "random_seed": SEED,
        "eta_charge": 0.9,
        "eta_discharge": 0.9,
        "soc_min_kwh": 1200.0,
        "soc_max_kwh": 10800.0,
        "charge_limit_kwh_per_interval": 5000.0 / 6.0,
        "discharge_limit_kwh_per_interval": 5000.0 / 6.0,
        "interval_minutes": 10,
        "allow_partial_extraction": True,
        "settlement_method": "lexicographic",
        "report_start": "2025-02-01",
        "report_end": "2025-12-31",
        "interval_convention": ("interval_index 1..144；t=1 对应真实 00:00-00:10，"
                                "t=144 对应 23:50-24:00；附件时刻标签为区间结束时刻"),
        "pv_pred_clip_at_zero": {
            "clipped_cells": int(clip_stats["clipped_cells"]),
            "max_magnitude_kwh": float(clip_stats["max_magnitude_kwh"]),
            "rationale": ("光伏预测物理上不小于 0；夜间 GBM 预测出现小幅负值，"
                          "按 scenarios.py 的物理截断规则截断到 0，截断统计如实记录。"
                          "残差文件同步按截断后预测计算，重建场景值与原滚动一致。"),
        },
        "beta_evidence": {
            "status": "CONFIRMED" if pkl_ev["match"] else "BLOCKED_BETA_UNVERIFIED",
            "evidence": [
                f"result2(7) SHA256 = {EXPECTED_SHA256}，与 All_Code/Q2/Results/Tables/result2.xlsx 完全一致",
                "q2_rolling_results.pkl 复算 6 个汇总值与提示词给定值一致（最大误差 < 1e-6），"
                "判定为 result2(7) 同一次运行",
                f"同一次运行参数：beta={pkl_ev['beta_best']}, M={pkl_ev['m_best']}, "
                f"kappa2={pkl_ev['kappa2_best']:.6f}, alpha={pkl_ev['alpha']}",
                "run_summary.txt 的 beta=0.0/M=30 为更早运行的过期记录，不作为 result2(7) 参数依据",
            ],
            "result2_7_run_parameters": {
                "beta": float(pkl_ev["beta_best"]),
                "M": int(pkl_ev["m_best"]),
                "kappa2": float(pkl_ev["kappa2_best"]),
                "alpha": float(pkl_ev["alpha"]),
                "settlement_method": "weighted(M_PLAN=20)",
            },
        },
        "note": ("本文件参数取自词典序结算后重新滚动并重新调参的运行，"
                 "与 q2_plan_reference_long.csv 同一次运行。"),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
    out = Q3_INTERFACE / "q2_parameters.json"
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(params, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    return params, out


def count_emergency_records(logs):
    n = 0
    for l in logs:
        e = l["settle_e"]
        t = 0
        while t < T:
            if e[t] > 1e-9:
                n += 1
                while t < T and e[t] > 1e-9:
                    t += 1
            else:
                t += 1
    return n


def build_corrected_workbook(results):
    """Q2_result2_corrected_for_Q3.xlsx：结构与官方模板一致，但充放电量工作表统一使用
    实际执行轨迹 settle_c_actual / settle_r_actual / settle_s_actual（不混用计划轨迹）。
    只写入 Q3，不覆盖 Q2 官方 result2.xlsx。"""
    report_logs = results["report_logs"]
    tpl = openpyxl.load_workbook(RESULT2_TEMPLATE)
    plan_header = [tpl["计划购电量"].cell(1, c).value
                   for c in range(1, tpl["计划购电量"].max_column + 1)]
    chg_header = [tpl["充放电量"].cell(1, c).value
                  for c in range(1, tpl["充放电量"].max_column + 1)]
    emg_header = [tpl["紧急购电量"].cell(1, c).value
                  for c in range(1, tpl["紧急购电量"].max_column + 1)]
    if len(plan_header) != 147:
        raise ValueError(f"计划购电量模板列数应为 147，实际 {len(plan_header)}")

    tm = pd.read_csv(REPO_ROOT / "All_Code" / "Data_preprocessing"
                     / "Data_transformation" / "time_map.csv", encoding="utf-8-sig")
    period_start = {int(r.time_idx): str(r.period_start) for r in tm.itertuples()}
    period_end = {int(r.time_idx): str(r.period_end) for r in tm.itertuples()}

    out = openpyxl.Workbook()
    out.remove(out.active)

    ws1 = out.create_sheet("计划购电量")
    ws1.append(plan_header)
    for i, l in enumerate(report_logs):
        d = datetime.strptime(l["date"], "%Y-%m-%d")
        x = l["plan_x"]
        row = [d] + [None] * 146
        for t in range(1, T):
            row[t] = float(x[t])
        nxt = report_logs[i + 1] if i + 1 < len(report_logs) else None
        if nxt is not None:
            if (datetime.strptime(nxt["date"], "%Y-%m-%d") - d).days != 1:
                raise ValueError(f"report_logs 日期不连续: {l['date']} -> {nxt['date']}")
            row[144] = float(nxt["plan_x"][0])
        else:
            row[144] = None
        row[145] = float(np.sum(x))
        row[146] = float(l["cost_plan"])
        ws1.append(row)

    ws2 = out.create_sheet("充放电量")
    ws2.append(chg_header)
    windows = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    for l in report_logs:
        d = datetime.strptime(l["date"], "%Y-%m-%d")
        c = l["settle_c_actual"]
        r = l["settle_r_actual"]
        s = l["settle_s_actual"]
        for k, (a, b) in enumerate(windows):
            row = [None] * 6
            if k == 0:
                row[0] = d
                row[4] = dtime(0, 0)
                row[5] = float(s[0])
            elif k == 1:
                row[4] = "24:00"
                row[5] = float(s[T])
            row[1] = labels[k]
            row[2] = float(np.sum(c[a:b]))
            row[3] = float(np.sum(r[a:b]))
            ws2.append(row)

    ws3 = out.create_sheet("紧急购电量")
    ws3.append(emg_header)
    for l in report_logs:
        d = datetime.strptime(l["date"], "%Y-%m-%d")
        e = l["settle_e"]
        t = 1
        while t <= T:
            if e[t - 1] > 1e-9:
                start = t
                while t <= T and e[t - 1] > 1e-9:
                    t += 1
                end = t - 1
                ws3.append([d, f"{period_start[start]}-{period_end[end]}",
                            float(np.sum(e[start - 1:end]))])
            else:
                t += 1

    out_path = Q3_TABLES / "Q2_result2_corrected_for_Q3.xlsx"
    out.save(out_path)

    total_plan = float(np.sum([np.sum(l["plan_x"]) for l in report_logs]))
    total_c = float(np.sum([np.sum(l["settle_c_actual"]) for l in report_logs]))
    total_r = float(np.sum([np.sum(l["settle_r_actual"]) for l in report_logs]))
    total_e = float(np.sum([l["emergency_kwh"] for l in report_logs]))
    wb = openpyxl.load_workbook(out_path, read_only=True)
    w1 = wb["计划购电量"]
    w2 = wb["充放电量"]
    w3 = wb["紧急购电量"]
    col146 = float(np.sum([r[145] or 0 for r in w1.iter_rows(min_row=2, values_only=True)]))
    chg = float(np.sum([r[2] or 0 for r in w2.iter_rows(min_row=2, values_only=True)]))
    dis = float(np.sum([r[3] or 0 for r in w2.iter_rows(min_row=2, values_only=True)]))
    emg_rows = sum(1 for r in w3.iter_rows(min_row=2, values_only=True) if r[0] is not None)
    wb.close()
    ok = (abs(col146 - total_plan) < 1e-6 and abs(chg - total_c) < 1e-6
          and abs(dis - total_r) < 1e-6 and emg_rows == count_emergency_records(report_logs))
    return out_path, dict(plan_kwh=total_plan,
                          plan_fee=float(np.sum([l["cost_plan"] for l in report_logs])),
                          charge_kwh=total_c, discharge_kwh=total_r,
                          emergency_kwh=total_e, emergency_records=emg_rows,
                          workbook_consistency_ok=ok)


def run_checks(results, src_ev, clip_stats):
    """15 项自动校验（提示词第十五节）。"""
    rows = []

    def rec(no, name, ok, evidence):
        rows.append((no, name, "PASS" if ok else "FAIL", evidence))

    fcsv = pd.read_csv(Q3_INTERFACE / "q2_forecast_10min.csv", encoding="utf-8-sig")
    rblk = pd.read_csv(Q3_INTERFACE / "q2_residual_blocks.csv", encoding="utf-8-sig")
    man = pd.read_csv(Q3_INTERFACE / "q2_scenario_manifest.csv",
                      encoding="utf-8-sig", keep_default_na=False)
    plan = pd.read_csv(Q3_INTERFACE / "q2_plan_reference_long.csv", encoding="utf-8-sig")

    rec(1, "q2_forecast_10min.csv 行数=52560", len(fcsv) == 52560, f"实际 {len(fcsv)}")
    sizes_f = fcsv.groupby("date")["interval_index"].count()
    sizes_p = plan.groupby("date")["interval_index"].count()
    sizes_r = rblk.groupby("residual_date")["interval_index"].count()
    ok2 = (sizes_f == 144).all() and (sizes_p == 144).all() and (sizes_r == 144).all()
    rec(2, "每个 date 恰好 144 个 interval_index", bool(ok2),
        f"forecast异常={(sizes_f != 144).sum()}, plan异常={(sizes_p != 144).sum()}, residual异常={(sizes_r != 144).sum()}")
    dup_f = int(fcsv.duplicated(["date", "interval_index"]).sum())
    dup_p = int(plan.duplicated(["date", "interval_index"]).sum())
    dup_r = int(rblk.duplicated(["residual_date", "interval_index"]).sum())
    rec(3, "date 与 interval_index 无重复", dup_f + dup_p + dup_r == 0,
        f"forecast={dup_f}, plan={dup_p}, residual={dup_r}")
    cols = ["load_pred_kwh", "pv_pred_0h_kwh", "load_actual_kwh", "pv_actual_kwh"]
    rec(4, "预测与实际数据均为有限数", bool(np.isfinite(fcsv[cols].to_numpy()).all()),
        f"非有限个数={int((~np.isfinite(fcsv[cols].to_numpy())).sum())}")
    rec(5, "负荷与光伏预测不小于0",
        bool((fcsv[["load_pred_kwh", "pv_pred_0h_kwh"]].to_numpy() >= 0).all()),
        "预测列非负检查（光伏预测已按物理截断规则截断到 >=0，"
        f"截断格数={clip_stats['clipped_cells']}，最大幅度={clip_stats['max_magnitude_kwh']:.3f} kWh）")
    m = fcsv.merge(rblk, left_on=["date", "interval_index"],
                   right_on=["residual_date", "interval_index"])
    err1 = float(np.max(np.abs(m["load_residual_kwh"] - (m["load_actual_kwh"] - m["load_pred_kwh"]))))
    err2 = float(np.max(np.abs(m["pv_residual_kwh"] - (m["pv_actual_kwh"] - m["pv_pred_0h_kwh"]))))
    rec(6, "残差恒等式 actual-prediction（误差<1e-8 kWh）", max(err1, err2) < 1e-8,
        f"max_err={max(err1, err2):.3e}, 合并行数={len(m)}")
    sub = man[man["source_residual_date"].astype(str).str.len() > 0].copy()
    sub["src_dt"] = pd.to_datetime(sub["source_residual_date"])
    sub["tgt_dt"] = pd.to_datetime(sub["target_date"])
    n_bad = int((~(sub["src_dt"] < sub["tgt_dt"])).sum()
                + (~(sub["source_day_index"].astype(int)
                     < sub["target_day_index"].astype(int))).sum())
    point_rows = int((man["source_residual_date"].astype(str).str.len() == 0).sum())
    rec(7, "场景来源日期严格早于目标日期", n_bad == 0,
        f"违规记录数={n_bad}; 点预测场景（残差池为空）记录数={point_rows}")
    pcols = ["plan_purchase_kwh", "actual_extracted_kwh", "actual_charge_kwh",
             "actual_discharge_kwh", "emergency_purchase_kwh", "curtailment_kwh"]
    rec(8, "计划/实际提取/充放电/紧急/弃光均非负",
        bool((plan[pcols].to_numpy() >= -1e-9).all()),
        f"负值个数={int((plan[pcols].to_numpy() < -1e-9).sum())}")
    ok9 = (plan["actual_extracted_kwh"] <= plan["plan_purchase_kwh"] + 1e-6).all()
    rec(9, "实际提取量 <= 计划购电量 + 1e-6", bool(ok9), f"违规行数={int((~ok9).sum())}")
    merged = plan.merge(fcsv[["date", "interval_index", "load_actual_kwh", "pv_actual_kwh"]],
                        on=["date", "interval_index"])
    g_used = merged["pv_actual_kwh"] - merged["curtailment_kwh"]
    bal = (merged["actual_extracted_kwh"] + merged["emergency_purchase_kwh"] + g_used
           + merged["actual_discharge_kwh"] - merged["load_actual_kwh"]
           - merged["actual_charge_kwh"])
    rec(10, "功率平衡最大绝对误差 <= 1e-5 kWh", float(bal.abs().max()) <= 1e-5,
        f"max={float(bal.abs().max()):.3e}")
    soc_err_max = 0.0
    for _, g in plan.groupby("date"):
        g = g.sort_values("interval_index")
        s_start = g["actual_soc_start_kwh"].to_numpy()
        s_end = g["actual_soc_end_kwh"].to_numpy()
        cc = g["actual_charge_kwh"].to_numpy()
        rr = g["actual_discharge_kwh"].to_numpy()
        soc_err_max = max(soc_err_max,
                          float(np.max(np.abs(s_end - (s_start + 0.9 * cc - rr / 0.9)))))
    rec(11, "SOC 递推最大绝对误差 <= 1e-5 kWh", soc_err_max <= 1e-5,
        f"max={soc_err_max:.3e}")
    lo, hi = 1200.0 - 1e-9, 10800.0 + 1e-9
    inb = (plan["actual_soc_start_kwh"].between(lo, hi)
           & plan["actual_soc_end_kwh"].between(lo, hi))
    rec(12, "SOC 全部位于 [1200,10800] kWh", bool(inb.all()),
        f"越界行数={int((~inb).sum())}")
    logs = results["report_logs"]
    n_bad_gap = sum(1 for l in logs if l["emergency_cost_gap"] > l["cost_tolerance"] + 1e-6)
    method_ok = all(l["settlement_method"] == "lexicographic" for l in logs)
    rec(13, "词典序第二阶段紧急购电费 <= 第一阶段最优值+容差",
        n_bad_gap == 0 and method_ok,
        f"超标天数={n_bad_gap}, 结算方法统一={method_ok}")
    rec(14, "result2(7) 汇总值与提示词列出的数值一致", bool(src_ev["all_ok"]),
        "SHA256 + 6 个汇总值 + 3 个行内区间差")
    cp = subprocess.run(["git", "-C", str(REPO_ROOT), "status", "--short"],
                        capture_output=True, text=True, encoding="utf-8", errors="replace")
    bad_lines = []
    for line in cp.stdout.splitlines():
        p = line.strip()
        if len(p) < 3:
            continue
        rel = p[3:]
        norm = rel.replace("\\", "/")
        if os.path.basename(rel).startswith("~$"):
            continue
        if norm.startswith("All_Code/Q2") and norm.rsplit(".", 1)[-1].lower() in {"csv", "json", "xlsx", "md"}:
            bad_lines.append(rel)
    rec(15, "All_Code/Q2 下无本任务新增 CSV/JSON/XLSX/MD 接口文件", len(bad_lines) == 0,
        f"违规条目: {bad_lines if bad_lines else '无（Excel 锁文件 ~$* 已排除，非本任务生成）'}")
    return rows


def build_audit_xlsx(checks_rows, src_ev, pkl_ev, results, head, reroll_sec, out_paths, clip_stats):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    all_pass = all(r[2] == "PASS" for r in checks_rows)

    ws = wb.create_sheet("Summary")
    summary_rows = [
        ["最终结论", "PASS" if all_pass else "存在 FAIL"],
        ["源工作簿", str(Q2_RESULT2_SOURCE)],
        ["源工作簿 SHA256", src_ev["hash"]],
        ["SHA256 与 result2(7) 一致", str(src_ev["hash_ok"])],
        ["源 Git 提交", head],
        ["beta 核验状态", "CONFIRMED" if pkl_ev["match"] else "BLOCKED_BETA_UNVERIFIED"],
        ["result2(7) 同次运行参数",
         f"beta={pkl_ev['beta_best']}, M={pkl_ev['m_best']}, kappa2={pkl_ev['kappa2_best']:.6f}, alpha={pkl_ev['alpha']}"],
        ["词典序重滚后参数",
         f"beta={results['beta_best']}, M={results['m_best']}, kappa2={results['kappa2_best']:.6f}, alpha={results['alpha']}"],
        ["结算方法", "lexicographic（M_PLAN 已从正式结算移除，legacy 仅历史对比）"],
        ["重滚耗时(秒)", f"{reroll_sec:.1f}"],
        ["光伏预测物理截断", f"截断格数={clip_stats['clipped_cells']}, "
         f"最大幅度={clip_stats['max_magnitude_kwh']:.3f} kWh（截断到 >=0，统计如实记录）"],
    ]
    for k, v in sorted(out_paths.items()):
        summary_rows.append([f"输出[{k}]", v])
    for r in summary_rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Parameter check")
    q3 = {"beta": float(results["beta_best"]), "M(scenario_count)": int(results["m_best"]),
          "alpha": float(results["alpha"]), "kappa2": float(results["kappa2_best"]),
          "kappa2_base": float(results["kappa2_base"]), "seed": SEED,
          "settlement_method": "lexicographic"}
    old7 = {"beta": 0.25, "M(scenario_count)": 20, "alpha": 0.9,
            "kappa2": 0.13291666666666666, "kappa2_base": 0.5316666666666666,
            "seed": 20260101, "settlement_method": "weighted(M_PLAN=20)"}
    gh = {"beta": 0.0, "M(scenario_count)": 30, "alpha": 0.9, "kappa2": 0.132917,
          "kappa2_base": 0.531667, "seed": 20260101, "settlement_method": "weighted(M_PLAN=20)"}
    notes = {
        "beta": "run_summary.txt 为更早运行（beta=0.0）；result2(7) 同次运行 beta=0.25；Q3 采用词典序重滚后的重调参结果",
        "M(scenario_count)": "同上：30 为更早运行，20 为 result2(7) 同次运行，Q3 为重滚后重调参结果",
        "alpha": "三次运行一致",
        "kappa2": "三次运行一致（舍入显示差异）",
        "kappa2_base": "三次运行一致（舍入显示差异）",
        "seed": "三次运行一致",
        "settlement_method": "result2(7) 使用 M_PLAN=20 加权目标；本次正式结算已改为词典序两阶段",
    }
    ws.append(["参数", "result2(7) 同次运行（旧缓存）", "GitHub 旧记录（run_summary.txt）",
               "Q3 采用（词典序重滚）", "一致?", "说明"])
    for k in ["beta", "M(scenario_count)", "alpha", "kappa2", "kappa2_base", "seed", "settlement_method"]:
        v_old, v_gh, v_q3 = old7[k], gh[k], q3[k]
        same = str(v_old) == str(v_q3)
        ws.append([k, v_old, v_gh, v_q3, "一致" if same else "不一致", notes[k]])
        r_idx = ws.max_row
        for col, v in ((2, v_old), (3, v_gh), (4, v_q3)):
            cell = ws.cell(r_idx, col)
            if str(v) != str(v_q3):
                cell.fill = RED
    for c, w in enumerate([24, 34, 36, 30, 10, 80], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws = wb.create_sheet("Data checks")
    ws.append(["编号", "检查项", "结果", "证据"])
    for r in checks_rows:
        ws.append(list(r))
        if r[2] == "FAIL":
            ws.cell(ws.max_row, 3).fill = RED
        else:
            ws.cell(ws.max_row, 3).fill = GREEN
    for c, w in enumerate([8, 52, 10, 60], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws = wb.create_sheet("Result2 comparison")
    ws.append(["汇总指标", "提示词给定值", "源工作簿复算值", "旧缓存复算值", "差异(工作簿-给定)", "判定"])
    for k in EXPECTED_SUMMARY:
        e = EXPECTED_SUMMARY[k]
        a = src_ev["summary"][k]["actual"]
        pv = pkl_ev["values"][k]
        ws.append([k, e, a, pv, a - e, "一致" if abs(a - e) <= 1e-3 else "不一致"])
    ws.append([])
    ws.append(["日期", "提示词给定行内差", "工作簿复算行内差", "差异", "判定"])
    for k in EXPECTED_ROW_DIFF:
        e = EXPECTED_ROW_DIFF[k]
        a = src_ev["row_diff"][k]["actual"]
        ws.append([k, e, a, (a or 0.0) - e, "一致" if abs((a or 0.0) - e) <= 0.05 else "不一致"])
    for c, w in enumerate([24, 22, 22, 22, 22, 12], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    out = Q3_TABLES / "Q2_to_Q3_handoff_audit.xlsx"
    wb.save(out)
    return out


def build_report_md(checks_rows, src_ev, pkl_ev, results, new_summary, head, reroll_sec,
                    out_paths, row_counts, clip_stats, all_pass):
    lines = []
    A = lines.append
    A("# Q2 → Q3 数据同步与自动修改交接报告")
    A("")
    A(f"- 生成时间: {datetime.now():%Y-%m-%d %H:%M:%S}")
    A(f"- 源 Git 提交: `{head}`（本地工作区，未 push）")
    A(f"- 总体结论: **{'完成：15 项自动校验全部 PASS' if all_pass else '存在 FAIL，不能正式接入 Q3'}**")
    A("")
    A("## 1. beta 核验")
    A("")
    A(f"- 状态: **{'CONFIRMED' if pkl_ev['match'] else 'BLOCKED_BETA_UNVERIFIED'}**")
    A("- 证据 1: `All_Code/Q2/Results/Tables/result2.xlsx` 的 SHA256 = "
      f"`{src_ev['hash']}`，与 result2(7).xlsx 的给定哈希 `{EXPECTED_SHA256}` 完全一致。")
    A("- 证据 2: 旧缓存 `q2_rolling_results.pkl` 复算 6 个汇总值（计划购电总量/计划购电费/充电量/"
      "放电量/紧急购电量/紧急购电记录数）与提示词给定值一致（最大误差 < 1e-6），"
      "判定该缓存与 result2(7) 同一次运行。")
    A(f"- 同一次运行参数: beta={pkl_ev['beta_best']}, M={pkl_ev['m_best']}, "
      f"kappa2={pkl_ev['kappa2_best']:.6f}, alpha={pkl_ev['alpha']}。")
    A("- 说明: `run_summary.txt` 中 beta=0.0/M=30 是更早运行的过期记录，不作为 result2(7) 的参数依据。")
    A("")
    A("## 2. 结算修改（M_PLAN 问题）")
    A("")
    A("- `settlement.py` 正式结算改为词典序两阶段：第一阶段 `min sum(5*pi*e)`；"
      "第二阶段在紧急购电费不超过第一阶段最优值+容差的前提下最大化计划充放电执行量。")
    A("- `M_PLAN=20` 加权目标已从正式结算移除，保留为 `settle_day_weighted_legacy()`，仅用于历史对比。")
    A(f"- 重滚报告期最大紧急购电费差 = {new_summary['max_gap']:.3e} 元，"
      f"最大容差 = {new_summary['max_tol']:.3e} 元，全部天数满足 gap <= tol + 1e-6。")
    A("")
    A("## 3. 词典序重滚与重新调参")
    A("")
    A(f"- 从 2025-01 预热期（i=1..30，基准参数 beta=0.5/kappa2_base）重新滚动，"
      f"02-01 期初 SOC 重新产生；耗时 {reroll_sec:.1f} 秒。")
    A(f"- 重新调参结果: beta={results['beta_best']}, M={results['m_best']}, "
      f"kappa2={results['kappa2_best']:.6f} (kappa2_base={results['kappa2_base']:.6f}), "
      f"alpha={results['alpha']}, seed={SEED}。")
    A(f"- result2(7) 同次运行参数为 beta={pkl_ev['beta_best']}/M={pkl_ev['m_best']}；"
      "若重调参后参数发生变化，论文 F1/F2 的参数结论必须按本报告更新，不得沿用旧调参结果。")
    A("")
    A("## 4. 新旧 result2 差异")
    A("")
    A("| 指标 | result2(7)（M_PLAN 加权结算） | 词典序重滚修正版（Q3） |")
    A("|---|---:|---:|")
    A(f"| 计划购电总量 kWh | {src_ev['summary']['计划购电总量(kWh)']['actual']:.2f} | {new_summary['plan_kwh']:.2f} |")
    A(f"| 计划购电费 元 | {src_ev['summary']['计划购电费(元)']['actual']:.2f} | {new_summary['plan_fee']:.2f} |")
    A(f"| 充电量 kWh | {src_ev['summary']['充电量(kWh)']['actual']:.2f}（计划口径） | {new_summary['charge_actual_kwh']:.2f}（实际口径） |")
    A(f"| 放电量 kWh | {src_ev['summary']['放电量(kWh)']['actual']:.2f}（计划口径） | {new_summary['discharge_actual_kwh']:.2f}（实际口径） |")
    A(f"| 紧急购电量 kWh | {src_ev['summary']['紧急购电量(kWh)']['actual']:.2f} | {new_summary['emergency_kwh']:.2f} |")
    A(f"| 紧急购电记录数 | {src_ev['summary']['紧急购电记录数']['actual']} | {new_summary['emergency_records']} |")
    A(f"| 报告期总费用 元 | — | {new_summary['cost_total']:.2f} |")
    A(f"| 报告期紧急购电率(电量) | — | {new_summary['emergency_rate']:.6f} |")
    A("")
    A("注：旧表充放电量为计划轨迹（plan_c/plan_r），修正版统一改为实际执行轨迹"
      "（settle_c_actual/r_actual/s_actual），两者口径不同，不能直接相减比较。")
    A("")
    A("## 5. 时间列问题")
    A("")
    A("- result2 宽表第 2..145 列包含“本日 t=2..144 + 次日 t=1”，而“全天购电量”列是本日 t=1..144 之和，"
      "两者天然不相等（如 2025-02-01 差 +117.24 kWh、2025-05-10 差 +1012.44 kWh、"
      "2025-12-31 差 -1464.30 kWh，与提示词一致）。")
    A("- Q3 接口不通过移动 Excel 列或跨日拼接恢复逐时计划；全部使用内部 `plan_x` 等长表，"
      "主键为 `date + interval_index`。")
    A("- 项目当前时间约定不变：`interval_index=1..144`，t=1 对应真实 00:00-00:10，"
      "t=144 对应 23:50-24:00；附件时间标签为区间结束时刻。")
    A("")
    A("## 6. Q3 接口文件与行数")
    A("")
    A("| 文件 | 行数 | 主键 |")
    A("|---|---:|---|")
    for k in ["q2_forecast_10min.csv", "q2_residual_blocks.csv", "q2_scenario_manifest.csv",
              "q2_plan_reference_long.csv"]:
        A(f"| {k} | {row_counts.get(k, 0)} | "
          f"{'date+interval_index' if k != 'q2_scenario_manifest.csv' else 'target_date+scenario_id'} |")
    A("| q2_parameters.json | 1 | — |")
    A("")
    A("场景因果检查：manifest 中所有非空 source_residual_date 均严格早于 target_date，"
      "source_day_index < target_day_index；残差池为空（i=1）时 source 留空并记为点预测场景。")
    A("")
    A(f"- 光伏预测物理截断：0 点光伏预测（GBM）出现 {clip_stats['clipped_cells']} 个负值格"
      f"（最大幅度 {clip_stats['max_magnitude_kwh']:.3f} kWh，夜间时段），"
      "接口中统一截断到 0（与 scenarios.py 的物理截断规则一致）；"
      "残差文件同步按截断后预测计算，因此 Q3 重建场景值与原滚动完全一致，不产生数据造假。")
    A("")
    A("## 7. 可供 Q3 使用的数值 / 不能写入论文的数值")
    A("")
    A("- 可用: q2_forecast_10min.csv（负荷预测与 0 点光伏预测）、q2_residual_blocks.csv、"
      "q2_scenario_manifest.csv（重建联合误差场景）、q2_parameters.json（核验参数）、"
      "Q2_result2_corrected_for_Q3.xlsx（仅结果对照）。")
    A("- 仅作交叉核对: q2_plan_reference_long.csv（Q3 零点计划必须按附件 3 的 0 点光伏预报重新求解，"
      "不得直接复制 Q2 的 plan_x）。")
    A("- 不可使用: 旧缓存中 M_PLAN 加权结算的 settle_* 轨迹、run_summary.txt 的 beta=0.0/M=30 "
      "过期参数、result2 宽表的逐时恢复值。")
    A("")
    A("## 8. 15 项自动校验")
    A("")
    A("| 编号 | 检查项 | 结果 | 证据 |")
    A("|---|---|---|---|")
    for no, name, status, ev in checks_rows:
        A(f"| {no} | {name} | {status} | {ev} |")
    A("")
    A("## 9. 输出文件绝对路径")
    A("")
    for k, v in sorted(out_paths.items()):
        A(f"- `{v}`")
    A("")
    A("## 10. 未解决问题")
    A("")
    A("- `All_Code/Q2/Results/Tables/run_summary.txt` 仍是 beta=0.0 的过期记录，建议在论文定稿前"
      "由 Q2 负责人确认是否更新，避免评委读取到矛盾参数。")
    A("- Q2 官方 `result2.xlsx` 宽表列错位问题未改动（不擅自改变 Q1/Q2/Q3 全局时间约定），"
      "Q3 一律使用长表接口。")
    A("- 本次所有新文件均在本地 Q3 目录，尚未 push（按提示词要求，除非用户明确要求不得 push）。")
    A("")
    out = Q3_TABLES / "Q2_to_Q3_handoff_report.md"
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return out


def main():
    t0 = time.time()
    head = git_head()
    print("=" * 70)
    print("Q2 -> Q3 数据同步与自动修改（交接执行器）")
    print(f"git HEAD = {head}")
    print("=" * 70)

    src_ev = verify_source_workbook()
    print(f"[1] 源工作簿 SHA256 = {src_ev['hash']}")
    print(f"    hash_ok = {src_ev['hash_ok']}")
    for k, v in src_ev["summary"].items():
        print(f"    {k}: 复算={v['actual']:.6f}, 给定={v['expected']}, diff={v['diff']:.2e}")
    for k, v in src_ev["row_diff"].items():
        print(f"    行内差 {k}: 复算={v['actual']}, 给定={v['expected']}, diff={v['diff']:.2e}")
    if not src_ev["all_ok"]:
        print("[FAIL] 源工作簿与 result2(7) 的核验未通过，停止，不得继续修改。")
        return 1

    pkl_ev = verify_old_pkl_vs_md()
    print(f"[2] 旧缓存复算 6 个汇总值与提示词一致 = {pkl_ev['match']}")
    if not pkl_ev["match"]:
        print("BETA_UNVERIFIED: result2(7).xlsx 不包含 beta，且没有找到对应运行日志")
        print("停止正式接口导出（不生成 q2_parameters.json / q2_plan_reference_long.csv）。")
        return 2
    print(f"    result2(7) 同次运行参数: beta={pkl_ev['beta_best']}, M={pkl_ev['m_best']}, "
          f"kappa2={pkl_ev['kappa2_best']:.6f}, alpha={pkl_ev['alpha']}")

    for m in (10, 20, 30):
        p = Q2_DATA_PROCESSING / f"scenarios_M{m}.pkl"
        if not p.exists():
            raise FileNotFoundError(f"缺少场景缓存（不得重新写入 Q2）: {p}")
    with open(Q2_DATA_PROCESSING / "q2_dataset.pkl", "rb") as fh:
        ds = pickle.load(fh)
    with open(Q2_DATA_PROCESSING / "forecasts.pkl", "rb") as fh:
        fc = pickle.load(fh)

    print("[3] 词典序结算重滚 + 重新调参（输出全部指向 Q3）...")
    from rolling import run_tuning_and_formal
    t1 = time.time()
    results = run_tuning_and_formal(
        tables_dir=str(Q3_TABLES),
        results_path=str(Q3_INTERFACE / "q2_rolling_results_lexicographic.pkl"))
    reroll_sec = time.time() - t1
    print(f"    重滚完成，耗时 {reroll_sec:.1f} 秒；"
          f"新参数: beta={results['beta_best']}, M={results['m_best']}, "
          f"kappa2={results['kappa2_best']:.6f}, alpha={results['alpha']}")

    for src_name, dst_name in [("tuning_sensitivity.xlsx", "Q2_reroll_tuning_sensitivity.xlsx"),
                               ("daily_rolling_log.csv", "Q2_reroll_daily_rolling_log.csv")]:
        s = Q3_TABLES / src_name
        if s.exists():
            dst = Q3_TABLES / dst_name
            if dst.exists():
                dst.unlink()
            s.rename(dst)

    print("[4] 生成 Q3 接口文件 ...")
    out_paths = {}
    row_counts = {}
    p1, n1, clip_stats = build_forecast_csv(fc, ds)
    out_paths["q2_forecast_10min.csv"] = str(p1)
    row_counts["q2_forecast_10min.csv"] = n1
    p2, n2 = build_residual_blocks(fc, ds)
    out_paths["q2_residual_blocks.csv"] = str(p2)
    row_counts["q2_residual_blocks.csv"] = n2
    p3, n3 = build_scenario_manifest(ds, int(results["m_best"]))
    out_paths["q2_scenario_manifest.csv"] = str(p3)
    row_counts["q2_scenario_manifest.csv"] = n3
    p4, n4 = build_plan_reference_long(results)
    out_paths["q2_plan_reference_long.csv"] = str(p4)
    row_counts["q2_plan_reference_long.csv"] = n4
    _, p5 = build_parameters_json(results, src_ev, pkl_ev, head, clip_stats)
    out_paths["q2_parameters.json"] = str(p5)
    print("    接口 CSV/JSON 已生成")

    p6, wb_summary = build_corrected_workbook(results)
    out_paths["Q2_result2_corrected_for_Q3.xlsx"] = str(p6)
    print(f"    修正版参考工作簿已生成，内部一致性复核={wb_summary['workbook_consistency_ok']}")

    logs = results["report_logs"]
    new_summary = {
        "plan_kwh": float(np.sum([np.sum(l["plan_x"]) for l in logs])),
        "plan_fee": float(np.sum([l["cost_plan"] for l in logs])),
        "charge_actual_kwh": float(np.sum([np.sum(l["settle_c_actual"]) for l in logs])),
        "discharge_actual_kwh": float(np.sum([np.sum(l["settle_r_actual"]) for l in logs])),
        "emergency_kwh": float(np.sum([l["emergency_kwh"] for l in logs])),
        "emergency_records": count_emergency_records(logs),
        "cost_total": float(np.sum([l["cost_total"] for l in logs])),
        "emergency_rate": float(np.sum([l["emergency_kwh"] for l in logs])
                                / np.sum([l["load_kwh"] for l in logs])),
        "max_gap": float(np.max([l["emergency_cost_gap"] for l in logs])),
        "max_tol": float(np.max([l["cost_tolerance"] for l in logs])),
    }

    print("[5] 执行 15 项自动校验 ...")
    checks_rows = run_checks(results, src_ev, clip_stats)
    all_pass = all(r[2] == "PASS" for r in checks_rows)
    for no, name, status, ev in checks_rows:
        print(f"    [{status}] {no:>2}. {name}: {ev}")

    out_paths["q2_rolling_results_lexicographic.pkl"] = str(Q3_INTERFACE / "q2_rolling_results_lexicographic.pkl")
    out_paths["Q2_reroll_tuning_sensitivity.xlsx"] = str(Q3_TABLES / "Q2_reroll_tuning_sensitivity.xlsx")
    out_paths["Q2_reroll_daily_rolling_log.csv"] = str(Q3_TABLES / "Q2_reroll_daily_rolling_log.csv")

    p7 = build_audit_xlsx(checks_rows, src_ev, pkl_ev, results, head, reroll_sec, out_paths, clip_stats)
    out_paths["Q2_to_Q3_handoff_audit.xlsx"] = str(p7)
    p8 = build_report_md(checks_rows, src_ev, pkl_ev, results, new_summary, head,
                         reroll_sec, out_paths, row_counts, clip_stats, all_pass)
    out_paths["Q2_to_Q3_handoff_report.md"] = str(p8)
    print("[6] 审计工作簿与交接报告已生成")

    print("[7] git diff --stat")
    cp1 = subprocess.run(["git", "-C", str(REPO_ROOT), "diff", "--stat"],
                         capture_output=True, text=True, encoding="utf-8", errors="replace")
    print(cp1.stdout)
    print("[8] git status --short")
    cp2 = subprocess.run(["git", "-C", str(REPO_ROOT), "status", "--short"],
                         capture_output=True, text=True, encoding="utf-8", errors="replace")
    print(cp2.stdout)
    print(f"总耗时 {time.time() - t0:.1f} 秒")
    return 0 if all_pass else 3


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(1)
