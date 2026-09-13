"""Reopen and independently audit the saved official-template Q4-3 workbook."""

from __future__ import annotations

import argparse
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent/"Results"/"Tables"   # 仓库统一结果目录
T=144
TOL=2e-5


def near(a,b,label):
    if a is None or b is None or abs(float(a)-float(b))>TOL:
        raise AssertionError(f"{label}: {a} != {b}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--strategy", required=True,
                        help="strategy used to create result4-3.xlsx")
    args = parser.parse_args()
    daily=pd.read_csv(ROOT/args.strategy/"daily.csv")
    slots=pd.read_csv(ROOT/args.strategy/"intervals.csv")
    book=openpyxl.load_workbook(ROOT/"result4-3.xlsx",read_only=False,data_only=True)
    expected_names=["计划购电量","调整购电量","充放电量","紧急购电量"]
    if book.sheetnames!=expected_names:
        raise AssertionError(f"wrong official-template sheets: {book.sheetnames}")
    dates=pd.to_datetime(daily.date).dt.date.tolist()
    x=slots.zero_plan_kwh.to_numpy().reshape(334,T)
    q=slots.final_plan_kwh.to_numpy().reshape(334,T)
    c=slots.charge_kwh.to_numpy().reshape(334,T)
    r=slots.discharge_kwh.to_numpy().reshape(334,T)
    e=slots.emergency_kwh.to_numpy().reshape(334,T)
    s0=slots.soc_start_kwh.to_numpy().reshape(334,T)[:,0]
    s1=slots.soc_end_kwh.to_numpy().reshape(334,T)[:,-1]
    for sheet_name,values,fee_column in (("计划购电量",x,"plan_cost_yuan"),
                                         ("调整购电量",q,"market_cost_yuan")):
        sheet=book[sheet_name]
        if sheet.max_row!=335 or sheet.max_column!=147:
            raise AssertionError(f"{sheet_name}: template extent mismatch")
        for i,row in enumerate(sheet.iter_rows(min_row=2,max_row=335,values_only=True)):
            if row[0].date()!=dates[i]:
                raise AssertionError(f"{sheet_name}: day {i} date mismatch")
            for t in range(1,T):
                near(row[t],values[i,t],f"{sheet_name}: day {i}, slot {t+1}")
            if i<333:
                near(row[144],values[i+1,0],f"{sheet_name}: next-day first slot {i}")
            elif row[144] is not None:
                raise AssertionError(f"{sheet_name}: final cross-day cell must be blank")
            near(row[145],np.sum(values[i]),f"{sheet_name}: day {i} natural sum")
            near(row[146],daily.loc[i,fee_column],f"{sheet_name}: day {i} fee")
    storage=book["充放电量"]
    if storage.max_row!=2005 or storage.max_column!=6:
        raise AssertionError("storage sheet extent mismatch")
    for i in range(334):
        for block in range(6):
            row=storage[i*6+block+2]
            if block==0 and row[0].value.date()!=dates[i]:
                raise AssertionError(f"storage day {i} date mismatch")
            near(row[2].value,np.sum(c[i,block*24:(block+1)*24]),f"charge day {i}, block {block}")
            near(row[3].value,np.sum(r[i,block*24:(block+1)*24]),f"discharge day {i}, block {block}")
            if block==0:
                near(row[5].value,s0[i],f"start SOC day {i}")
            if block==1:
                near(row[5].value,s1[i],f"end SOC day {i}")
    emergency=book["紧急购电量"]
    event_sum=sum(float(row[2] or 0) for row in emergency.iter_rows(min_row=2,values_only=True))
    near(event_sum,np.sum(e),"all emergency events")
    _ev=0
    for _i in range(334):
        _pos=e[_i]>1e-8
        _ev+=int(np.flatnonzero(_pos & ~np.r_[False,_pos[:-1]]).size)
    if emergency.max_row!=_ev+1:
        raise AssertionError(f"emergency event count mismatch: {emergency.max_row-1} != {_ev}")
    print(f"PASS result4-3.xlsx ({args.strategy}): 334 dates, 96,192 purchase cells, 2,004 storage rows, "
          f"{emergency.max_row-1:,} emergency events, full cross-day mapping and cost reconciliation")


if __name__=="__main__":
    main()
