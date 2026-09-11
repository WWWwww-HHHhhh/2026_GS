# -*- coding: utf-8 -*-
"""Q2'' 官方模板导出 result2.xlsx（计划购电量 / 充放电量 / 紧急购电量）。"""
import os, pickle
from datetime import datetime, time as dtime
import numpy as np
import openpyxl

GS = r"E:\2.University_materials\4.University_life\6.Study_materials\others\freshman_year(secong semester)\2026_GS"
Q2P = os.path.join(GS, "Q2''")
DATA_PROC = os.path.join(Q2P, "Data_processing")
TABLES = os.path.join(Q2P, "Results", "Tables")
TEMPLATE = os.path.join(GS, "Data", "附件", "附件5", "result2.xlsx")
OUT = os.path.join(TABLES, "result2.xlsx")
T = 144

def _hm(minutes):
    h, m = divmod(int(round(minutes)), 60)
    return f"{h}:{m:02d}"
def period_start(t):
    return _hm((t-1)*10)
def period_end(t):
    return "0:00+1" if t*10 == 1440 else _hm(t*10)

def main():
    res = pickle.load(open(os.path.join(DATA_PROC, "q2_rolling_results.pkl"), "rb"))
    report = res["report_logs"]
    tpl = openpyxl.load_workbook(TEMPLATE)
    plan_tpl = tpl["计划购电量"]; chg_tpl = tpl["充放电量"]; emg_tpl = tpl["紧急购电量"]
    plan_header = [plan_tpl.cell(1, c).value for c in range(1, plan_tpl.max_column + 1)]
    chg_header = [chg_tpl.cell(1, c).value for c in range(1, chg_tpl.max_column + 1)]
    emg_header = [emg_tpl.cell(1, c).value for c in range(1, emg_tpl.max_column + 1)]
    if len(plan_header) != 147:
        raise ValueError(f"计划购电量模板列数应为147，实际{len(plan_header)}")

    out = openpyxl.Workbook(); out.remove(out.active)
    ws1 = out.create_sheet("计划购电量"); ws1.append(plan_header)
    for i, l in enumerate(report):
        d = datetime.strptime(l["date"], "%Y-%m-%d")
        x = l["plan_x"]
        row = [d] + [None]*146
        for t in range(1, T):
            row[t] = float(x[t])
        nxt = report[i+1] if i+1 < len(report) else None
        if nxt is not None:
            row[144] = float(nxt["plan_x"][0])
        else:
            row[144] = None
        row[145] = float(np.sum(x))
        row[146] = float(l["cost_plan"])
        ws1.append(row)

    ws2 = out.create_sheet("充放电量"); ws2.append(chg_header)
    for l in report:
        d = datetime.strptime(l["date"], "%Y-%m-%d")
        c = l["plan_c"]; r = l["plan_r"]; s = l["plan_s"]
        windows = [(0,24),(24,48),(48,72),(72,96),(96,120),(120,144)]
        labels = ["0:00-4:00","4:00-8:00","8:00-12:00","12:00-16:00","16:00-20:00","20:00-24:00"]
        for k,(a,b) in enumerate(windows):
            row = [None]*6
            if k == 0:
                row[0] = d; row[4] = dtime(0,0); row[5] = float(s[0])
            elif k == 1:
                row[4] = "24:00"; row[5] = float(s[T])
            row[1] = labels[k]; row[2] = float(np.sum(c[a:b])); row[3] = float(np.sum(r[a:b]))
            ws2.append(row)

    ws3 = out.create_sheet("紧急购电量"); ws3.append(emg_header)
    for l in report:
        d = datetime.strptime(l["date"], "%Y-%m-%d")
        e = l["settle_e"]
        t = 1
        while t <= T:
            if e[t-1] > 1e-9:
                start = t
                while t <= T and e[t-1] > 1e-9:
                    t += 1
                end = t - 1
                ws3.append([d, f"{period_start(start)}-{period_end(end)}", float(np.sum(e[start-1:end]))])
            else:
                t += 1
    out.save(OUT)
    print("result2.xlsx saved", OUT)

if __name__ == "__main__":
    main()
