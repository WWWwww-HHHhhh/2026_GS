from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter


def as_date(value):
    """把载荷中的 'YYYY-MM-DD' 字符串转成真日期对象，保证 Excel 日期格式与排序正常。"""
    if value is None or isinstance(value, date):
        return value
    if isinstance(value, datetime):
        return value.date()
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except ValueError:
        return value

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = Path(r"D:\MathModel\2026_GS\Data\附件\附件5\result3.xlsx")

TIME_LABELS = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]


def write_value_block(ws, payload_rows: list[list], ncols: int) -> None:
    """把 334×ncols 的数值写入第 2 行起（第 1 列日期不写入，保持模板列结构）。"""
    for i, row in enumerate(payload_rows):
        excel_row = i + 2
        for j, value in enumerate(row[1:1 + ncols]):
            if value is None:
                continue
            ws.cell(row=excel_row, column=j + 2, value=value)


def build(folder: Path, template: Path, out_path: Path) -> dict:
    payload = json.loads((folder / "result3_payload.json").read_text(encoding="utf-8"))
    days = int(payload["report_days"])
    assert days == 334, f"载荷天数应为 334，实际 {days}"
    assert len(payload["plan"]) == days and len(payload["adjusted"]) == days
    assert len(payload["charge"]) == days * 6, "充放电量应为 334×6 行"

    wb = openpyxl.load_workbook(template)

    # 1) 计划购电量 / 调整购电量
    for sheet_name, rows in (("计划购电量", payload["plan"]), ("调整购电量", payload["adjusted"])):
        ws = wb[sheet_name]
        write_value_block(ws, rows, ncols=146)
        for r in range(2, days + 2):
            for c in range(2, 146):
                ws.cell(row=r, column=c).number_format = "0.000000"
            ws.cell(row=r, column=146).number_format = "#,##0.00"
            ws.cell(row=r, column=147).number_format = "#,##0.00"
        ws.freeze_panes = "B2"

    # 2) 充放电量（每日 6 行；日期只在当日首行给出，时间段标签逐行写入）
    ws = wb["充放电量"]
    for d in range(days):
        base = 2 + d * 6
        for k in range(6):
            row = payload["charge"][d * 6 + k]
            # 注意：openpyxl 的 cell(value=None) 不会清空旧值，必须对 .value 赋值
            c_date = ws.cell(row=base + k, column=1)
            c_date.value = as_date(row[0])
            c_date.number_format = "yyyy-mm-dd"
            ws.cell(row=base + k, column=2).value = row[1]
            for col_idx, value in ((3, row[2]), (4, row[3]), (5, row[4]), (6, row[5])):
                cell = ws.cell(row=base + k, column=col_idx)
                cell.value = value
            ws.cell(row=base + k, column=3).number_format = "#,##0.000000"
            ws.cell(row=base + k, column=4).number_format = "#,##0.000000"
            ws.cell(row=base + k, column=6).number_format = "#,##0.000000"
    ws.freeze_panes = "A2"

    # 3) 紧急购电量（清空模板示例行后写入事件明细）
    ws = wb["紧急购电量"]
    for r in range(2, 12):
        for c in range(1, 4):
            ws.cell(row=r, column=c).value = None
    thin = Border(*[Side(style="thin", color="9CA3AF")] * 4)
    for i, row in enumerate(payload["emergency"]):
        excel_row = i + 2
        ws.cell(row=excel_row, column=1, value=as_date(row[0])).number_format = "yyyy-mm-dd"
        ws.cell(row=excel_row, column=2, value=row[1])
        ws.cell(row=excel_row, column=3, value=row[2]).number_format = "#,##0.000000"
        for c in range(1, 4):
            ws.cell(row=excel_row, column=c).border = thin
    ws.freeze_panes = "A2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {
        "output": str(out_path),
        "plan_rows": len(payload["plan"]),
        "adjusted_rows": len(payload["adjusted"]),
        "storage_rows": len(payload["charge"]),
        "emergency_rows": len(payload["emergency"]),
        "sheets": wb.sheetnames,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--folder", type=Path, default=HERE.parent / "Results" / "Tables")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args()
    out = args.out or (args.folder / "result3.xlsx")
    info = build(args.folder, args.template, out)
    print(json.dumps(info, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
